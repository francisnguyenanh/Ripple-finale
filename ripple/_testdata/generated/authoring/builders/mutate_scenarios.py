"""Scenario mutation and materialization engine for Harbor Procurement Mini (Stage 07)."""

import shutil
from pathlib import Path
import openpyxl
import yaml


def materialize_scenario(
    base_bundle_dir: Path,
    target_bundle_dir: Path,
    scenario_id: str,
) -> None:
    scenario_id = scenario_id.upper()
    print(f"Materializing scenario {scenario_id} into {target_bundle_dir}...")

    # Copy base bundle
    if target_bundle_dir.exists():
        shutil.rmtree(target_bundle_dir)
    shutil.copytree(base_bundle_dir, target_bundle_dir)

    cfg_file = target_bundle_dir / "config" / "approval.yaml"
    excel_file = target_bundle_dir / "data" / "approval_matrix.xlsx"
    emergency_file = target_bundle_dir / "app" / "emergency.py"
    fin_guide_file = target_bundle_dir / "docs" / "finance_desk_guide.md"

    if scenario_id == "S0":
        # S0 is already baseline
        pass

    elif scenario_id == "S1":
        # Standard config: 500,000 JPY
        with open(cfg_file, "r", encoding="utf-8") as f:
            cfg = yaml.safe_load(f)
        cfg["routing_rules"]["standard_manager_threshold_jpy"] = 500_000
        with open(cfg_file, "w", encoding="utf-8") as f:
            yaml.safe_dump(cfg, f)

        # Emergency code follows 500,000 JPY
        content = emergency_file.read_text(encoding="utf-8")
        content = content.replace("1_000_000", "500_000")
        emergency_file.write_text(content, encoding="utf-8")

        # Excel is already 500k in S0 -> NO_OP

    elif scenario_id == "S2":
        # Standard config: 750,000 JPY
        with open(cfg_file, "r", encoding="utf-8") as f:
            cfg = yaml.safe_load(f)
        cfg["routing_rules"]["standard_manager_threshold_jpy"] = 750_000
        with open(cfg_file, "w", encoding="utf-8") as f:
            yaml.safe_dump(cfg, f)

        # Emergency code follows 750,000 JPY
        content = emergency_file.read_text(encoding="utf-8")
        content = content.replace("1_000_000", "750_000")
        emergency_file.write_text(content, encoding="utf-8")

        # Excel updated to 750,000 JPY
        wb = openpyxl.load_workbook(excel_file)
        ws = wb["RoutingRules"]
        ws["B3"] = 750_000  # Manager row is row 3 in 1-based indexing (or cell B3)
        wb.save(excel_file)

    elif scenario_id == "S4":
        # From S1 (500k config), Excel drifts to 650k
        materialize_scenario(base_bundle_dir, target_bundle_dir, "S1")
        wb = openpyxl.load_workbook(excel_file)
        ws = wb["RoutingRules"]
        ws["B3"] = 650_000
        wb.save(excel_file)

    elif scenario_id == "S5":
        # From S1, finance desk guide is corrected to CAPEX only
        materialize_scenario(base_bundle_dir, target_bundle_dir, "S1")
        content = fin_guide_file.read_text(encoding="utf-8")
        content = content.replace(
            "all purchasing commitments of **2,000,000 JPY or higher**",
            "capital expenditures (CAPEX) of **2,000,000 JPY or higher**",
        )
        fin_guide_file.write_text(content, encoding="utf-8")

    print(f"Scenario {scenario_id} materialized successfully.")


if __name__ == "__main__":
    base = Path(__file__).resolve().parent.parent.parent
    base_b = base / "bundles" / "harbor_demo" / "input"
    s1_b = base / "bundles" / "harbor_demo_s1" / "input"
    materialize_scenario(base_b, s1_b, "S1")
