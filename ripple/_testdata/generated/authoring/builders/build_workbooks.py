"""Deterministic Excel workbook generator for Harbor Procurement Mini using openpyxl."""

import csv
import json
from datetime import datetime
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def build_approval_matrix(output_path: Path, manager_threshold_jpy: int = 500_000) -> None:
    """Builds data/approval_matrix.xlsx.
    
    In S0 baseline, manager_threshold_jpy is set to 500,000 JPY (C01 value conflict).
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RoutingRules"
    ws.views.sheetView[0].showGridLines = True

    # Styling
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    headers = ["Role", "Threshold_JPY", "Condition", "Notes"]
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    rows = [
        ("AUTO", 499_999, "Standard OPEX below Manager limit", "Automated system clearance"),
        ("MANAGER", manager_threshold_jpy, "Standard / General limit", "Operational guidance limit"),
        ("FINANCE", 2_000_000, "CAPEX purchases >= 2,000,000 JPY", "Capex committee review"),
        ("DIRECTOR", 0, "High Risk Vendor (Any amount)", "Mandatory executive approval"),
        ("REJECTED", 0, "Blocked Vendor / Invalid Input", "Immediate rejection"),
    ]

    for r_idx, r_data in enumerate(rows, start=2):
        role_c = ws.cell(row=r_idx, column=1, value=r_data[0])
        thresh_c = ws.cell(row=r_idx, column=2, value=r_data[1])
        cond_c = ws.cell(row=r_idx, column=3, value=r_data[2])
        notes_c = ws.cell(row=r_idx, column=4, value=r_data[3])

        role_c.font = data_font
        role_c.alignment = align_center
        role_c.border = thin_border

        thresh_c.font = data_font
        thresh_c.alignment = align_right
        thresh_c.number_format = "#,##0"
        thresh_c.border = thin_border

        cond_c.font = data_font
        cond_c.alignment = align_left
        cond_c.border = thin_border

        notes_c.font = data_font
        notes_c.alignment = align_left
        notes_c.border = thin_border

    # Unrelated operational sheet to test preservation of untouched content
    ws2 = wb.create_sheet(title="DepartmentCodes")
    ws2.append(["Code", "Department", "CostCenter"])
    ws2.cell(row=1, column=1).font = header_font
    ws2.cell(row=1, column=1).fill = header_fill
    ws2.cell(row=1, column=2).font = header_font
    ws2.cell(row=1, column=2).fill = header_fill
    ws2.cell(row=1, column=3).font = header_font
    ws2.cell(row=1, column=3).fill = header_fill

    dept_rows = [
        ("PROC", "Procurement", "CC-101"),
        ("OPS", "Operations", "CC-102"),
        ("FIN", "Finance", "CC-103"),
        ("EXEC", "Executive", "CC-104"),
    ]
    for d_idx, d_data in enumerate(dept_rows, start=2):
        ws2.append(list(d_data))

    # Auto-adjust column widths
    for sheet in (ws, ws2):
        for col in sheet.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Freeze header pane
    ws.freeze_panes = "A2"
    ws2.freeze_panes = "A2"

    wb.save(output_path)


def build_supplier_directory(csv_path: Path, output_path: Path) -> None:
    """Builds data/supplier_directory.xlsx from suppliers.csv."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Vendors"
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    headers = ["Supplier_ID", "Vendor_Name", "Rating", "Status", "Payment_Terms"]
    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col_num)
        c.fill = header_fill
        c.font = header_font

    with open(csv_path, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for r_idx, row in enumerate(reader, start=2):
            rating = "HIGH_RISK" if row["risk"] == "HIGH" else "STANDARD"
            ws.append([
                row["supplier_id"],
                row["display_name"],
                rating,
                row["status"],
                "Net 30 Days",
            ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    ws.freeze_panes = "A2"
    wb.save(output_path)


def build_procurement_summary(invoices_jsonl: Path, output_path: Path) -> None:
    """Builds data/procurement_summary_202609.xlsx from invoices.jsonl."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MonthlySummary"
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    headers = ["Metric", "Value_JPY", "Count"]
    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col_num)
        c.fill = header_fill
        c.font = header_font

    total_net = 0
    total_tax = 0
    total_freight = 0
    total_gross = 0
    count = 0

    with open(invoices_jsonl, "r", encoding="utf-8") as f:
        for line in f:
            if not line.strip():
                continue
            data = json.loads(line)
            total_net += data["goods_net_jpy"]
            total_tax += data["tax_jpy"]
            total_freight += data["freight_jpy"]
            total_gross += data["gross_jpy"]
            count += 1

    summary_rows = [
        ("Total Goods Net (JPY)", total_net, count),
        ("Total Estimated Tax (JPY)", total_tax, count),
        ("Total Freight (JPY)", total_freight, count),
        ("Total Gross Billed (JPY)", total_gross, count),
    ]

    for row in summary_rows:
        ws.append(list(row))

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 6, 16)

    ws.freeze_panes = "A2"
    wb.save(output_path)


def build_all_workbooks(data_dir: Path) -> None:
    matrix_path = data_dir / "approval_matrix.xlsx"
    suppliers_csv = data_dir / "suppliers.csv"
    supplier_wb = data_dir / "supplier_directory.xlsx"
    inv_jsonl = data_dir / "invoices.jsonl"
    summary_wb = data_dir / "procurement_summary_202609.xlsx"

    print("Building approval_matrix.xlsx with C01 (500,000 JPY manager threshold)...")
    build_approval_matrix(matrix_path, manager_threshold_jpy=500_000)

    print("Building supplier_directory.xlsx...")
    build_supplier_directory(suppliers_csv, supplier_wb)

    print("Building procurement_summary_202609.xlsx...")
    build_procurement_summary(inv_jsonl, summary_wb)
    print("All workbooks successfully built!")


if __name__ == "__main__":
    base = Path(__file__).resolve().parent.parent.parent
    target_data = base / "bundles" / "harbor_demo" / "input" / "data"
    build_all_workbooks(target_data)
