"""Independent verification and knowledge rebuild from the new source.

Two deterministic, zero-provider-call executors and the post-publish consumer
that feeds them.

  * `consume_source_published` -- called by `job_worker.claim_and_admit_one`
    for a claimed `SOURCE_PUBLISHED` outbox event whose payload carries
    `kind == "apply"` (the durable link `ApplyExecutor`'s publish writes in
    the SAME transaction as the head CAS). It creates -- through the existing
    `store.create_job`, deterministic ids, so a duplicate delivery or a crash
    between "create" and "ack" is a no-op -- the two jobs below, THEN acks.

  * kind `verification` (`VerificationExecutor`): builds the reviewed
    `core.change_intelligence.v2.schema.VerificationReport` from the ACTUAL
    published blob bytes of every artifact of the published revision (never
    from the apply job's recorded hashes): checksum, diff against the parent
    revision (expected / actual / missing / unexpected), parse of the changed
    artifacts, the recorded typed edits re-read from the bytes, an INDEPENDENT
    acceptance oracle (env `RIPPLE_VERIFICATION_ORACLE_FILE`, tighten-only:
    absent => `acceptance_tests` NOT_RUN => the report can never be PASSED),
    and the approved plan's manual work carried verbatim (unproven required
    manual work => PARTIAL). The report is the job's immutable
    `execution_result` plus additive keys and a `report_digest`.

  * kind `km_rebuild` (`KmRebuildExecutor`): re-ingests the artifacts of the
    published revision (NEW observations under the NEW revision id, the old
    revision's rows are never touched), rebuilds the retrieval index, routes
    the candidates extracted from changed artifacts through the EXISTING
    `invalidation` executor (decision id `apply-<apply_job_id>`), schedules
    reanalysis as existing `extraction` jobs (scope FULL when dependency
    coverage is missing) and records carry-forward (a decision is carried only
    when its evidence binding validates on the new source), counterevidence
    from the NEW index and a KM CANDIDATE id. It never writes the workspace
    head, a knowledge revision or a semantic decision; a head that moved past
    the published revision makes the result `SUPERSEDED` and schedules nothing.

An additional, operator-supplied acceptance source: the verification contract named
by the approved REVIEW (`verification_contract` in its result), resolved through the immutable
apply -> approval -> review chain, mapped to the published artifact by `logical_path` and
evaluated against the READ-BACK published bytes (`_resolve_contract` / `_evaluate_contract`).
It is merged tighten-only with the env oracle: a contract that is missing/unreadable keeps
`acceptance_tests` NOT_RUN, a mismatch FAILs, and the env oracle can only add a failure.

`verification` and `km_rebuild` are deliberately absent from
`JOB_CREATE_PERMISSION` (a user / model cannot mint a report through the public
job route); retry goes through the existing retry route under
`change.accept_result`.
"""

from __future__ import annotations

import hashlib
import io
import json
import os

from core.change_intelligence.v2.serialize import canonical_hash
from core.change_intelligence.v2.source_ingest_adapters import PARSER_VERSION, parser_for
from infra.adapters.v2_postgres import (
    BlobIntegrityError, StaleVersion, StoreConflict, TenantScopeError,
)
from infra.adapters.v2_restore import load_restore_approval
from infra.worker import job_quota
from infra.worker.apply import _pending_and_scope, _resolve_job
from infra.worker.budget_config import BudgetExceeded
from infra.worker.execution_types import (
    ExecutionContext, ExecutionOutcome, ExecutionRequest, NonRetryableExecutionError,
    ProviderCallBudget,
)

VERIFICATION_KIND = "verification"
REBUILD_KIND = "km_rebuild"
ORACLE_ENV = "RIPPLE_VERIFICATION_ORACLE_FILE"
REPORT_SCHEMA_VERSION = "1.0.0"
REQUIRED_CHECKS = ("published_bytes_readback", "expected_edits", "parse_after",
                   "acceptance_tests")
_MISSING = object()


#: Job-id prefixes this batch derives deterministically. The public create
#: route refuses them (`infra/v2_policy_routes.py::create_job`), so a user
#: cannot squat a stage id; `ensure_job` still turns a foreign row into a
#: named `StageIdConflict` instead of an escaped exception.
#: `ingest-` (the ingest job of an import/refresh commit) and `refresh-`
#: (the decision id the refresh consumer stamps on its stale markers / causation).
RESERVED_JOB_ID_PREFIXES = ("verify-", "rebuild-", "inval-", "reanalysis-", "prop-revise-",
                            "ingest-", "refresh-")


class StageIdConflict(StoreConflict):
    """A job with a stage-derived id already exists with another kind."""


def verification_job_id(apply_job_id: str) -> str:
    return f"verify-{apply_job_id}"


def rebuild_job_id(apply_job_id: str) -> str:
    return f"rebuild-{apply_job_id}"


def apply_decision_id(apply_job_id: str) -> str:
    """The decision id the rebuild routes its invalidation jobs under."""
    return f"apply-{apply_job_id}"


def knowledge_candidate_id(apply_job_id: str) -> str:
    return f"km-cand-{apply_job_id}"


# ----------------------------------------------------------------------------- consumer

def ensure_job(store, tenant_id: str, workspace_id: str, job_id: str, kind: str,
               payload: dict, *, reserve: bool = False) -> bool:
    """Idempotently create a job through the existing `create_job` primitive.
    Deterministic ids make a duplicate delivery / crash-before-ack a no-op. A
    replay under a MOVED head cannot use `create_job`'s own idempotency key
    (its request hash includes `expected_version`), so existence is checked
    first. Returns True only when this call created it. `reserve=True` (a provider kind created by the system) reserves the job's worst case
    against the project cap inside `create_job`'s transaction and raises `BudgetExceeded` instead of creating the job when it does not fit."""
    for _ in range(5):
        try:
            existing = store.job_record(tenant_id, workspace_id, job_id)
        except TenantScopeError:
            existing = None
        if existing is not None:
            if existing["kind"] != kind:
                raise StageIdConflict(f"job {job_id!r} exists with kind {existing['kind']!r}")
            return False
        head = store.head(tenant_id, workspace_id)
        try:
            reservation = job_quota.reservation_for_system_job(kind) if reserve else None
            store.create_job(tenant_id, workspace_id, job_id, kind, payload,
                             expected_version=head.version,
                             idempotency_key=f"job-create-{job_id}",
                             budget_reservation=reservation)
            return True
        except (StaleVersion, StoreConflict):
            continue  # head moved / concurrent create: re-check existence
    raise StoreConflict(f"could not ensure job {job_id!r}")


def consume_source_published(store, tenant_id: str, workspace_id: str, claim,
                             worker_id: str, ack) -> bool:
    """Create the verification + rebuild jobs for a published apply, THEN ack
    (create-then-ack: a crash between the two redelivers the event and the
    deterministic ids make the second pass a no-op). Returns True if acked."""
    payload = claim.payload if isinstance(claim.payload, dict) else {}
    apply_job_id = payload.get("job_id")
    published = payload.get("source_revision_id")
    if not (isinstance(apply_job_id, str) and apply_job_id
            and isinstance(published, str) and published):
        return False  # malformed: leave it visible (bounded redelivery), create nothing
    base = {"apply_job_id": apply_job_id, "published_revision_id": published,
            "approval_job_id": payload.get("approval_job_id"),
            "target_artifact_id": payload.get("target_artifact_id")}
    conflicts = []
    for job_id, kind in ((verification_job_id(apply_job_id), VERIFICATION_KIND),
                         (rebuild_job_id(apply_job_id), REBUILD_KIND)):
        try:
            ensure_job(store, tenant_id, workspace_id, job_id, kind, dict(base))
        except StoreConflict as exc:
            conflicts.append({"job_id": job_id, "kind": kind, "error": str(exc)})
    if conflicts:
        # Fail closed and visible: the stage never runs (no report => never
        # eligible), an audit row names the squatted ids, and the event is
        # acked so the queue is not starved by a redelivery loop.
        from psycopg2.extras import Json
        with store._tx() as cur:
            cur.execute("""INSERT INTO v2_audit (tenant_id, workspace_id, event_type, actor_id, payload)
                VALUES (%s,%s,'STAGE_JOB_ID_CONFLICT','service',%s)""",
                        (tenant_id, workspace_id,
                         Json({"apply_job_id": apply_job_id, "conflicts": conflicts})))
    return ack(store, tenant_id, workspace_id, claim, worker_id)


# ----------------------------------------------------------------------------- helpers

def _json_safe(value):
    return json.loads(json.dumps(value, default=str))


def _scope_mismatch_section(ctx: ExecutionContext) -> dict:
    """The report keeps the durable, read-only scope_mismatch records of the workspace's
    completed finding jobs (same value, containing scopes = NOT a conflict). Not a finding, no lifecycle, never a blocker."""
    job = ctx.job
    return ctx.store.list_scope_comparisons(job["tenant_id"], job["workspace_id"])


def _payload_ids(ctx: ExecutionContext) -> tuple[str, str]:
    payload = ctx.job["payload"]
    apply_job_id = payload.get("apply_job_id")
    published = payload.get("published_revision_id")
    for key, value in (("apply_job_id", apply_job_id), ("published_revision_id", published)):
        if not isinstance(value, str) or not value:
            raise NonRetryableExecutionError(f"{key}_required", f"payload.{key}")
    return apply_job_id, published


def _revision_view(ctx, revision_id: str) -> dict:
    try:
        return ctx.store.source_revision_view(
            ctx.job["tenant_id"], ctx.job["workspace_id"], revision_id)
    except TenantScopeError as exc:
        raise NonRetryableExecutionError("published_revision_missing", revision_id) from exc


def _probe(store, artifact: dict) -> tuple[dict, bytes | None]:
    """Read the ACTUAL stored bytes of one artifact straight from the blob
    port (not `_verified_blob`, which would hide a corrupt blob's real hash).
    Returns the report entry and the bytes only when the checksum verifies."""
    digest = artifact["blob_hash"]
    entry = {"blob_hash": digest, "sha256": None, "readable": False, "checksum_ok": False,
             "byte_size": None}
    try:
        content = store.blobs.read(digest)
    except (KeyError, FileNotFoundError):
        return entry, None
    if not isinstance(content, bytes):
        return entry, None
    hex_digest = hashlib.sha256(content).hexdigest()
    entry.update(sha256=hex_digest, readable=True, byte_size=len(content),
                 checksum_ok=(f"sha256:{hex_digest}" == digest))
    return entry, (content if entry["checksum_ok"] else None)


def _artifact_name(artifact: dict) -> str:
    return artifact.get("logical_path") or artifact["artifact_id"]


def _kind_of(name: str, edit_kind: str | None, expected: bool) -> str | None:
    lower = name.lower()
    if lower.endswith((".yaml", ".yml")):
        return "yaml"
    if lower.endswith(".xlsx"):
        return "xlsx"
    if expected:
        return {"yaml_typed_edit": "yaml", "xlsx_typed_edit": "xlsx"}.get(edit_kind)
    return None


def _parse(kind: str, content: bytes):
    """Independent parse of the actual bytes. Returns (ok, detail, parsed)."""
    if kind == "yaml":
        import yaml
        try:
            return True, "yaml parsed", yaml.safe_load(content.decode("utf-8"))
        except Exception as exc:  # noqa: BLE001 - any parse failure is a FAIL, not a crash
            return False, f"{type(exc).__name__}: {exc}", None
    import openpyxl
    try:
        return True, "xlsx parsed", openpyxl.load_workbook(io.BytesIO(content))
    except Exception as exc:  # noqa: BLE001
        return False, f"{type(exc).__name__}: {exc}", None


def _yaml_lookup(data, path: str):
    node = data
    for key in path.split("."):
        if not isinstance(node, dict) or key not in node:
            return _MISSING
        node = node[key]
    return node


def _edit_locator(edit_kind: str | None, edit: dict) -> str:
    if edit_kind == "xlsx_typed_edit":
        return f"{edit.get('sheet')}!{edit.get('cell')}"
    return str(edit.get("path"))


def _actual_value(kind: str, parsed, edit_kind: str | None, edit: dict):
    if kind == "yaml":
        return _yaml_lookup(parsed, str(edit.get("path")))
    sheet, cell = edit.get("sheet"), edit.get("cell")
    if sheet not in parsed.sheetnames:
        return _MISSING
    return parsed[sheet][cell].value


def _strict_eq(a, b) -> bool:
    """Type-strict structural equality: `True != 1`, but `30 == 30.0` (a
    number-typed cell may round-trip int/float); lists/dicts recursively."""
    if isinstance(a, bool) or isinstance(b, bool):
        return isinstance(a, bool) and isinstance(b, bool) and a == b
    if isinstance(a, dict) and isinstance(b, dict):
        return a.keys() == b.keys() and all(_strict_eq(a[k], b[k]) for k in a)
    if isinstance(a, (list, tuple)) and isinstance(b, (list, tuple)):
        return len(a) == len(b) and all(_strict_eq(x, y) for x, y in zip(a, b))
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return a == b
    return type(a) is type(b) and a == b


def _tree_diff(expected, actual, path=""):
    """Every difference between two parsed yaml trees (added / removed /
    changed), each with its locator."""
    out = []
    if isinstance(expected, dict) and isinstance(actual, dict):
        for key in sorted(set(expected) | set(actual), key=str):
            loc = f"{path}.{key}" if path else str(key)
            if key not in actual:
                out.append({"locator": loc, "kind": "removed",
                            "expected": _json_safe(expected[key]), "actual": None})
            elif key not in expected:
                out.append({"locator": loc, "kind": "added", "expected": None,
                            "actual": _json_safe(actual[key])})
            else:
                out.extend(_tree_diff(expected[key], actual[key], loc))
    elif not _strict_eq(expected, actual):
        out.append({"locator": path or "$", "kind": "changed",
                    "expected": _json_safe(expected), "actual": _json_safe(actual)})
    return out


def _xlsx_snapshot(workbook) -> dict:
    """Semantic content of a workbook: ordered sheet names, every non-empty
    cell value and defined names (named ranges). Serialization-only details
    (zip timestamps, style ids, doc properties) are the declared normalization
    of an `openpyxl` re-save (apply policy) and are not compared."""
    cells = {}
    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cells[f"{ws.title}!{cell.coordinate}"] = cell.value
    names = {}
    dn = workbook.defined_names
    items = dn.items() if hasattr(dn, "items") else [(d.name, d) for d in dn.definedName]
    for name, defined in items:
        names[str(name)] = str(getattr(defined, "attr_text", defined))
    return {"sheets": list(workbook.sheetnames), "cells": cells, "names": names}


def _content_diff(store, edit_kind, before_hash, recorded_edits, actual_kind, actual_doc):
    """Compare the FULL published content of a target against the expected
    after-state = the parent (before) bytes + the approved recorded edits ONLY,
    built with the same typed-edit functions the apply gate uses. Returns the
    list of differences (empty when only the approved edits landed)."""
    from infra.worker.proposal import _xlsx_typed_edits, _yaml_typed_edits
    _entry, before = _probe(store, {"blob_hash": before_hash})
    if before is None:
        return [{"locator": "$", "kind": "expected_state_unbuildable",
                 "detail": "parent (before) bytes unreadable or checksum mismatch"}]
    try:
        if edit_kind == "yaml_typed_edit" and actual_kind == "yaml":
            import yaml
            edits_in = [{"path": e.get("path"), "expected_current": e.get("before"),
                         "proposed": e.get("after")} for e in recorded_edits]
            _a, _n, text = _yaml_typed_edits(before.decode("utf-8"), edits_in)
            return _tree_diff(yaml.safe_load(text) or {}, actual_doc or {})
        if edit_kind == "xlsx_typed_edit" and actual_kind == "xlsx":
            import openpyxl
            edits_in = [{"sheet": e.get("sheet"), "cell": e.get("cell"),
                         "expected_current": e.get("before"), "proposed": e.get("after")}
                        for e in recorded_edits]
            _a, _n, new_bytes = _xlsx_typed_edits(before, edits_in)
            exp = _xlsx_snapshot(openpyxl.load_workbook(io.BytesIO(new_bytes)))
            act = _xlsx_snapshot(actual_doc)
            out = []
            if exp["sheets"] != act["sheets"]:
                out.append({"locator": "$sheets", "kind": "changed",
                            "expected": exp["sheets"], "actual": act["sheets"]})
            for key in sorted(set(exp["cells"]) | set(act["cells"])):
                e, a = exp["cells"].get(key, _MISSING), act["cells"].get(key, _MISSING)
                if e is _MISSING or a is _MISSING or not _strict_eq(e, a):
                    out.append({"locator": key,
                                "kind": ("added" if e is _MISSING else
                                         "removed" if a is _MISSING else "changed"),
                                "expected": None if e is _MISSING else _json_safe(e),
                                "actual": None if a is _MISSING else _json_safe(a)})
            for key in sorted(set(exp["names"]) | set(act["names"])):
                if exp["names"].get(key) != act["names"].get(key):
                    out.append({"locator": f"$name:{key}", "kind": "changed",
                                "expected": exp["names"].get(key),
                                "actual": act["names"].get(key)})
            return out
    except NonRetryableExecutionError as exc:
        return [{"locator": "$", "kind": "expected_state_unbuildable", "detail": str(exc)}]
    return []


PIN_ENV = "RIPPLE_VERIFICATION_ORACLE_SHA256"


def _load_oracles():
    """Returns `(oracles, meta, problem)`.

    * no env => `(None, None, None)` (acceptance NOT_RUN);
    * missing / unreadable file => `(None, None, "oracle_unavailable")`;
    * the operator pin `RIPPLE_VERIFICATION_ORACLE_SHA256` is MANDATORY: an
      oracle file with no pin (unset or blank) => `(None, meta, "oracle_pin_missing")`,
      and a pin that differs from the file's sha256 => `(None, meta, "oracle_hash_mismatch")`;
    * all of these keep `acceptance_tests` NOT_RUN (completion stays blocked: the
      completion gate refuses `machine_verification_not_pass`).
    A readable file with malformed content is the named terminal failure
    `oracle_config_invalid`, never a PASS. The file is operator-provided and
    reviewed; its sha256/version are pinned in the report and its content is
    never handed to a model (this module makes no provider call)."""
    path = os.environ.get(ORACLE_ENV)
    if not path:
        return None, None, None
    try:
        with open(path, "rb") as handle:
            data = handle.read()
    except OSError:
        return None, None, "oracle_unavailable"
    try:
        raw = json.loads(data.decode("utf-8"))
    except ValueError as exc:
        raise NonRetryableExecutionError("oracle_config_invalid", f"{type(exc).__name__}: {exc}")
    oracles = raw.get("oracles") if isinstance(raw, dict) else None
    if not isinstance(oracles, list):
        raise NonRetryableExecutionError("oracle_config_invalid", "'oracles' list required")
    for item in oracles:
        kind = item.get("kind") if isinstance(item, dict) else None
        keys = {"yaml_value_equals": ("oracle_id", "artifact_id", "path"),
                "xlsx_cell_equals": ("oracle_id", "artifact_id", "sheet", "cell")}.get(kind)
        if not (keys and all(isinstance(item.get(k), str) and item.get(k) for k in keys) and "expected" in item):
            raise NonRetryableExecutionError("oracle_config_invalid", f"bad oracle {item!r}")
        if kind == "xlsx_cell_equals":
            # normalise to the row shape the report already uses ("path" = Sheet!Cell); same mandatory hash pin
            item["path"] = f"{item['sheet']}!{item['cell'].upper()}"
    version = raw.get("version")
    meta = {"sha256": hashlib.sha256(data).hexdigest(),
            "version": version if isinstance(version, (str, int)) else "unversioned",
            "oracle_ids": [o["oracle_id"] for o in oracles]}
    pin = (os.environ.get(PIN_ENV) or "").strip().lower()
    if not pin:
        # An unpinned oracle is unverifiable; it can never make acceptance PASS.
        return None, meta, "oracle_pin_missing"
    if pin != meta["sha256"]:
        return None, meta, "oracle_hash_mismatch"
    return oracles, meta, None


def oracle_actual(doc, oracle):
    """The value an oracle reads from a parsed artifact, or _MISSING (wrong kind of document, absent sheet/path/
    cell, empty cell). The oracle kind must match the document kind: a yaml oracle never reads a workbook."""
    if doc is None:
        return _MISSING
    if doc[0] == "yaml" and oracle.get("kind") == "yaml_value_equals":
        return _yaml_lookup(doc[1], oracle["path"])
    if doc[0] == "xlsx" and oracle.get("kind") == "xlsx_cell_equals":
        wb = doc[1]
        if oracle["sheet"] not in wb.sheetnames:
            return _MISSING
        try:
            value = wb[oracle["sheet"]][oracle["cell"].upper()].value
        except (ValueError, KeyError):
            return _MISSING
        return _MISSING if value is None else value
    return _MISSING


def _resolve_contract(ctx: ExecutionContext, approval_result: dict):
    """The operator-supplied verification contract the approved REVIEW named, resolved
    through the immutable apply -> approval -> review chain (never "latest", never an env/repo
    file). Returns `(named, contract, problem, embedded_sha)`: `named` is False only when the
    review result carries no contract at all (the legacy, contract-less flow)."""
    review_job_id = approval_result.get("review_job_id")
    try:
        review_job = _resolve_job(ctx, review_job_id, expected_kind="review", label="review")
    except NonRetryableExecutionError:
        return False, None, None, None
    info = (review_job.get("execution_result") or {}).get("verification_contract")
    if not isinstance(info, dict) or not info.get("contract_id"):
        return False, None, None, None
    embedded_sha = info.get("contract_sha256")
    try:
        contract = ctx.store.verification_contract(
            ctx.job["tenant_id"], ctx.job["workspace_id"], info["contract_id"])
    except Exception as exc:  # noqa: BLE001 -- unreadable => NOT_RUN, never a PASS
        return True, None, f"contract_unreadable:{type(exc).__name__}", embedded_sha
    if contract is None:
        return True, None, "contract_missing", embedded_sha
    return True, contract, None, embedded_sha


def _evaluate_contract(contract, problem, embedded_sha, approval_result, artifacts,
                       parsed_docs, expected_ids):
    """Evaluate the contract against the READ-BACK published bytes (already checksum-verified and
    parsed by `execute`), mapped to the published artifact by `logical_path`. Returns
    `(status, detail, rows)` with status PASS / FAIL / NOT_RUN: missing / unreadable => NOT_RUN,
    any mismatch (hash, proposal, artifact, kind, value) => FAIL; never a hardcoded PASS."""
    from infra.worker import verification_contract as vc
    if contract is None:
        return "NOT_RUN", f"{problem}: the named verification contract could not be read", []
    row = {"contract_id": contract.get("contract_id"),
           "contract_sha256": contract.get("contract_sha256"),
           "logical_path": contract.get("logical_path"), "locator": contract.get("locator"),
           "kind": contract.get("kind"), "expected": _json_safe(contract.get("expected")),
           "author_account_id": contract.get("author_account_id"),
           "provenance": contract.get("provenance") or {}, "actual": None, "status": "FAIL"}
    if (vc.contract_hash_of(contract) != contract.get("contract_sha256")
            or contract.get("contract_sha256") != embedded_sha):
        return "FAIL", "contract_hash_mismatch: stored contract differs from the one the review bound", [row]
    if contract.get("proposal_job_id") != approval_result.get("proposal_job_id"):
        return "FAIL", "contract_proposal_mismatch: contract belongs to another proposal", [row]
    matches = [aid for aid, art in artifacts.items()
               if _artifact_name(art) == contract.get("logical_path")]
    if not matches:
        row["status"] = "NOT_RUN"
        return "NOT_RUN", "published revision has no artifact at the contract's logical_path", [row]
    artifact_id = matches[0]
    if artifact_id not in expected_ids:
        return "FAIL", "contract maps to a different artifact than the approved target", [row]
    doc = parsed_docs.get(artifact_id)
    if doc is None:
        row["status"] = "NOT_RUN"
        return "NOT_RUN", "published bytes unreadable or unparseable", [row]
    if {"yaml_value_equals": "yaml", "xlsx_cell_equals": "xlsx"}.get(contract.get("kind")) != doc[0]:
        return "FAIL", "contract kind does not match the published artifact type", [row]
    status, actual = vc.evaluate_value(contract, doc)
    row.update(status=status, actual=actual, artifact_id=artifact_id)
    if status != "PASS":
        return "FAIL", (f"contract disagrees with the published bytes at "
                        f"{contract.get('locator')}"), [row]
    return "PASS", "operator-supplied contract agrees with the read-back published bytes", [row]


class VerificationExecutor:
    """See module docstring. Zero provider calls (`external = False`)."""

    external = False
    executor_id = "verification.independent_report.v1"

    def prepare(self, ctx: ExecutionContext) -> ExecutionRequest:
        _payload_ids(ctx)
        return ExecutionRequest(prompt=None,
                                schema={"type": "object", "executor": self.executor_id},
                                provider_calls_planned=0, external=False)

    def execute(self, ctx: ExecutionContext, request: ExecutionRequest,
                calls: ProviderCallBudget) -> ExecutionOutcome:
        job = ctx.job
        tenant_id, workspace_id = job["tenant_id"], job["workspace_id"]
        apply_job_id, published = _payload_ids(ctx)

        # Approved chain (immutable, COMPLETED upstream jobs). The apply job's
        # own terminal row is NOT required: the durable event is written inside
        # the publish transaction, before the apply job is finalized.
        try:
            apply_job = ctx.store.job_record(tenant_id, workspace_id, apply_job_id)
        except TenantScopeError as exc:
            raise NonRetryableExecutionError("apply_job_not_found", apply_job_id) from exc
        if apply_job.get("kind") != "apply":
            raise NonRetryableExecutionError("apply_job_wrong_kind", apply_job_id)
        if (apply_job.get("payload") or {}).get("restore_approval_id") is not None:
            return self._execute_restore(ctx, apply_job, apply_job_id, published)
        approval_job_id = (apply_job.get("payload") or {}).get("approval_job_id")
        approval_job = _resolve_job(ctx, approval_job_id, expected_kind="approval",
                                    label="approval")
        approval_result = approval_job.get("execution_result") or {}
        proposal_job = _resolve_job(ctx, approval_result.get("proposal_job_id"),
                                    expected_kind="proposal", label="proposal")
        proposal_result = proposal_job.get("execution_result") or {}
        target = proposal_result.get("target_artifact_id")
        edit_kind = proposal_result.get("edit_kind")
        recorded_edits = [dict(e) for e in (proposal_result.get("edits") or [])
                          if isinstance(e, dict)]
        if not isinstance(target, str) or not target:
            raise NonRetryableExecutionError("proposal_missing_target_artifact_id",
                                             str(approval_result.get("proposal_job_id")))

        view = _revision_view(ctx, published)
        parent_id = view["parent_revision_id"]
        parent_map = {}
        if parent_id:
            parent_map = {a["artifact_id"]: a["blob_hash"]
                          for a in _revision_view(ctx, parent_id)["artifacts"]}
        artifacts = {a["artifact_id"]: a for a in view["artifacts"]}
        expected_ids = [target]

        # ---- 1. actual bytes / checksum of EVERY artifact of the published revision
        entries: dict[str, dict] = {}
        contents: dict[str, bytes | None] = {}
        for artifact_id, art in artifacts.items():
            entries[artifact_id], contents[artifact_id] = _probe(ctx.store, art)
            entries[artifact_id]["changed"] = parent_map.get(artifact_id) != art["blob_hash"]
        unreadable = sorted(a for a, e in entries.items()
                            if not (e["readable"] and e["checksum_ok"]))

        # ---- 2. expected / actual / unexpected / missing
        changed = sorted(a for a, e in entries.items() if e["changed"])
        removed = sorted(a for a in parent_map if a not in artifacts)
        missing = [a for a in expected_ids if a not in changed]
        unexpected = sorted([a for a in changed if a not in expected_ids] + removed)
        expected_edits = [{"artifact_id": a, "edit_kind": edit_kind,
                           "edits": [{"locator": _edit_locator(edit_kind, e),
                                      "before": e.get("before"), "after": e.get("after")}
                                     for e in recorded_edits]} for a in expected_ids]
        actual_edits = [{"artifact_id": a, "before_blob_hash": parent_map.get(a),
                         "after_blob_hash": artifacts[a]["blob_hash"],
                         "sha256": entries[a]["sha256"]} for a in changed]
        unexpected_edits = [{"artifact_id": a, "before_blob_hash": parent_map.get(a),
                             "after_blob_hash": (artifacts[a]["blob_hash"]
                                                 if a in artifacts else None),
                             "removed": a not in artifacts} for a in unexpected]
        missing_edits = [{"artifact_id": a,
                          "reason": ("artifact_absent_from_published_revision"
                                     if a not in artifacts
                                     else "published_blob_equals_parent_blob")}
                         for a in missing]

        # ---- 3. parse the actual bytes (expected + every changed artifact)
        parse_targets = sorted(set(expected_ids) | set(changed))
        parsed_docs: dict[str, tuple] = {}
        parse_problems, parse_notes = [], []
        for artifact_id in parse_targets:
            art = artifacts.get(artifact_id)
            if art is None:
                parse_problems.append(f"{artifact_id}: absent from published revision")
                continue
            kind = _kind_of(_artifact_name(art), edit_kind, artifact_id in expected_ids)
            content = contents.get(artifact_id)
            if content is None:
                parse_problems.append(f"{artifact_id}: unreadable or checksum mismatch")
                continue
            if kind is None:
                parse_notes.append(f"{artifact_id}: no parser (not checked)")
                continue
            ok, detail, parsed = _parse(kind, content)
            if ok:
                parsed_docs[artifact_id] = (kind, parsed)
            else:
                parse_problems.append(f"{artifact_id}: {detail}")

        # ---- 4. values READ from the actual bytes; recorded edits re-checked
        actual_values: dict[str, dict] = {}
        edit_problems = []
        for artifact_id in expected_ids:
            if artifact_id not in parsed_docs:
                continue
            kind, parsed = parsed_docs[artifact_id]
            values = {}
            for edit in recorded_edits:
                value = _actual_value(kind, parsed, edit_kind, edit)
                locator = _edit_locator(edit_kind, edit)
                values[locator] = None if value is _MISSING else _json_safe(value)
                if value is _MISSING or not _strict_eq(value, edit.get("after")):
                    edit_problems.append(
                        f"{artifact_id}:{locator} expected {edit.get('after')!r}, "
                        f"actual {None if value is _MISSING else value!r}")
            actual_values[artifact_id] = values
        for artifact_id in expected_ids:
            if artifact_id in missing:
                edit_problems.append(f"{artifact_id}: expected edit never landed")
            elif artifact_id not in parsed_docs:
                edit_problems.append(f"{artifact_id}: actual bytes cannot be read/parsed")
        if not recorded_edits:
            edit_problems.append("approved proposal recorded no typed edit to verify")

        # ---- 4b. unexpected write INSIDE an expected artifact (R-1): the FULL
        #          content vs (before bytes + approved edits only)
        intra: dict[str, list] = {}
        for artifact_id in expected_ids:
            if (artifact_id in parsed_docs and artifact_id in changed
                    and parent_map.get(artifact_id) and recorded_edits):
                kind, parsed = parsed_docs[artifact_id]
                diffs = _content_diff(ctx.store, edit_kind, parent_map[artifact_id],
                                      recorded_edits, kind, parsed)
                if diffs:
                    intra[artifact_id] = diffs
        if intra:
            unexpected = sorted(set(unexpected) | set(intra))
            unexpected_edits = [{"artifact_id": a, "before_blob_hash": parent_map.get(a),
                                 "after_blob_hash": (artifacts[a]["blob_hash"]
                                                     if a in artifacts else None),
                                 "sha256": entries.get(a, {}).get("sha256"),
                                 "removed": a not in artifacts,
                                 "scope": "content" if a in intra else "artifact",
                                 "differences": intra.get(a, [])} for a in unexpected]

        # ---- 5. independent acceptance oracle
        oracles, oracle_meta, oracle_problem = _load_oracles()
        relevant = [o for o in (oracles or []) if o["artifact_id"] in expected_ids]
        oracle_rows = []
        if oracle_problem:
            acceptance = ("NOT_RUN", f"{oracle_problem}: the independent oracle did not run")
        elif not relevant:
            acceptance = ("NOT_RUN", "no independent oracle covers the changed artifact "
                          "(configure RIPPLE_VERIFICATION_ORACLE_FILE)")
        else:
            failures = []
            for oracle in relevant:
                doc = parsed_docs.get(oracle["artifact_id"])
                actual = oracle_actual(doc, oracle)
                passed = actual is not _MISSING and _strict_eq(actual, oracle["expected"])
                oracle_rows.append({
                    "oracle_id": oracle["oracle_id"], "artifact_id": oracle["artifact_id"],
                    "path": oracle["path"], "expected": _json_safe(oracle["expected"]),
                    "actual": None if actual is _MISSING else _json_safe(actual),
                    "requirement_ref": oracle.get("requirement_ref"),
                    "status": "PASS" if passed else "FAIL"})
                if not passed:
                    failures.append(oracle["oracle_id"])
            acceptance = (("FAIL", f"oracle(s) disagree: {failures}") if failures
                          else ("PASS", f"{len(relevant)} independent oracle(s) agree"))

        # ---- 5b. Operator-supplied verification contract (named by the approved
        #          review). Tighten-only merge with the optional env oracle above: a FAIL on
        #          either side fails, a NOT_RUN on either side (contract unreadable, env oracle
        #          configured but unusable) keeps the report from PASSING, and the env oracle can
        #          never turn a contract FAIL/NOT_RUN into a PASS. Without a contract named by
        #          the review the legacy env-oracle-only behaviour is unchanged.
        contract_named, contract, contract_problem, contract_sha = _resolve_contract(
            ctx, approval_result)
        contract_rows = []
        if contract_named:
            c_status, c_detail, contract_rows = _evaluate_contract(
                contract, contract_problem, contract_sha, approval_result, artifacts,
                parsed_docs, expected_ids)
            parts = [(c_status, f"contract: {c_detail}")]
            if oracle_problem or relevant:
                parts.append((acceptance[0], f"env oracle: {acceptance[1]}"))
            statuses = [p[0] for p in parts]
            merged = ("FAIL" if "FAIL" in statuses else "NOT_RUN" if "NOT_RUN" in statuses
                      else "PASS")
            acceptance = (merged, "; ".join(p[1] for p in parts))

        # ---- 6. manual work carried verbatim from the approved plan
        before_hash = parent_map.get(target)
        after_hash = artifacts[target]["blob_hash"] if target in artifacts else None
        pending, scope = _pending_and_scope(
            ctx, approval_result, proposal_result, target, before_hash, after_hash, published)
        manual_work = [{"index": i, "action": e.get("action"), "reason": e.get("reason"),
                        "proof": None} for i, e in enumerate(pending)]
        open_required_work = [f"manual_work[{m['index']}]: {m['reason']}" for m in manual_work]

        # ---- checks
        def _check(status, detail, required=True):
            return {"status": status, "required": required, "detail": detail}

        checks = {
            "published_bytes_readback": _check(
                "FAIL" if unreadable else "PASS",
                (f"unreadable: {unreadable}" if unreadable
                 else f"{len(entries)} artifact(s) read back and checksum-verified")),
            "expected_edits": _check(
                "FAIL" if edit_problems else "PASS",
                "; ".join(edit_problems) or "recorded edits present in the actual bytes"),
            "parse_after": _check(
                "FAIL" if parse_problems else "PASS",
                "; ".join(parse_problems + parse_notes) or "actual bytes parse"),
            "acceptance_tests": _check(acceptance[0], acceptance[1]),
            "simulation_exact_after": _check(
                "SKIPPED_NOT_APPLICABLE" if edit_kind in ("yaml_typed_edit", "xlsx_typed_edit")
                else "NOT_RUN",
                "no executable logic in a yaml/xlsx snapshot"
                if edit_kind in ("yaml_typed_edit", "xlsx_typed_edit")
                else f"edit_kind {edit_kind!r} needs a sandbox simulation",
                required=edit_kind not in ("yaml_typed_edit", "xlsx_typed_edit")),
        }

        reasons = ([f"missing_write:{a}" for a in missing]
                   + [f"unexpected_write:{a}" for a in unexpected]
                   + [f"unreadable_artifact:{a}" for a in unreadable])
        failed = bool(missing or unexpected or unreadable)
        not_run = False
        for name, chk in checks.items():
            if chk["required"] and chk["status"] == "FAIL":
                failed = True
                reasons.append(f"required_check_failed:{name}")
            elif chk["required"] and chk["status"] == "NOT_RUN":
                not_run = True
                reasons.append(f"required_check_not_run:{name}")
        if oracle_problem:
            reasons.append(oracle_problem)
        if open_required_work:
            reasons.append("manual_action_unproven")
        status = ("FAILED" if failed else "NOT_RUN" if not_run
                  else "PARTIAL" if open_required_work else "PASSED")

        body = {
            "id": job["job_id"],
            "schema_version": REPORT_SCHEMA_VERSION,
            "apply_job_id": apply_job_id,
            "approval_job_id": approval_job_id,
            "parent_revision_id": parent_id,
            "expected_edits": expected_edits,
            "actual_edits": actual_edits,
            "unexpected_edits": unexpected_edits,
            "missing_edits": missing_edits,
            "parsing_result": {"status": checks["parse_after"]["status"],
                               "problems": parse_problems, "notes": parse_notes},
            "tests_result": {"status": checks["acceptance_tests"]["status"],
                             "oracles": oracle_rows, "oracle_config": oracle_meta,
                             **({"contracts": contract_rows} if contract_named else {})},
            "oracle_config": oracle_meta,
            "simulation_result": {"status": checks["simulation_exact_after"]["status"],
                                  "detail": checks["simulation_exact_after"]["detail"]},
            "evidence_result": {"status": checks["published_bytes_readback"]["status"],
                                "artifacts_checked": len(entries), "unreadable": unreadable},
            "open_required_work": open_required_work,
            "scope": _json_safe({**scope, "parent_revision_id": parent_id}),
            "verification_scope": _json_safe(scope),
            "coverage": {"artifacts_total": len(entries),
                         "artifacts_readable": len(entries) - len(unreadable),
                         "changed": changed, "expected": expected_ids},
            "after_source_revision_id": published,
            "after_knowledge_revision_id": None,
            "status": status,
            "reasons": reasons,
            "checks": checks,
            "artifacts": {a: entries[a] for a in sorted(entries)},
            "actual_values": actual_values,
            "manual_work": manual_work,
            "scope_mismatches": _scope_mismatch_section(ctx),
        }
        body = _json_safe(body)
        result = {**body, "report_digest": canonical_hash(body)}
        return ExecutionOutcome(result=result, model_id=None)

    def _execute_restore(self, ctx: ExecutionContext, apply_job: dict, apply_job_id: str,
                         published: str) -> ExecutionOutcome:
        """Verify a RESTORE publication. The expected after-state is
        the WHOLE target revision (every artifact, hash AND bytes, read
        independently from the target's own rows), not a typed edit; removals
        are part of it. The oracle contract is unchanged:
        no oracle covering a touched artifact => `acceptance_tests` NOT_RUN =>
        the report can never be PASSED."""
        store = ctx.store
        tenant_id, workspace_id = ctx.job["tenant_id"], ctx.job["workspace_id"]
        approval = load_restore_approval(
            store, tenant_id, workspace_id,
            (apply_job.get("payload") or {}).get("restore_approval_id"))
        if approval is None:
            raise NonRetryableExecutionError(
                "restore_approval_not_found", str(apply_job_id))
        target_id = approval["target_revision_id"]
        view = _revision_view(ctx, published)
        target_view = _revision_view(ctx, target_id)
        parent_id = view["parent_revision_id"]
        parent_map = {}
        if parent_id:
            parent_map = {a["artifact_id"]: a["blob_hash"]
                          for a in _revision_view(ctx, parent_id)["artifacts"]}
        artifacts = {a["artifact_id"]: a for a in view["artifacts"]}
        target_arts = {a["artifact_id"]: a for a in target_view["artifacts"]}
        target_map = {a: art["blob_hash"] for a, art in target_arts.items()}
        expected_ids = sorted(a for a in set(parent_map) | set(target_map)
                              if parent_map.get(a) != target_map.get(a))

        # ---- 1. actual bytes / checksum of EVERY artifact of the published revision
        entries: dict[str, dict] = {}
        contents: dict[str, bytes | None] = {}
        for artifact_id, art in artifacts.items():
            entries[artifact_id], contents[artifact_id] = _probe(store, art)
            entries[artifact_id]["changed"] = parent_map.get(artifact_id) != art["blob_hash"]
        unreadable = sorted(a for a, e in entries.items()
                            if not (e["readable"] and e["checksum_ok"]))

        # ---- 2. whole-content compare against the TARGET (hash AND bytes)
        target_contents = {a: _probe(store, art)[1] for a, art in target_arts.items()}
        content_problems = []
        for artifact_id in sorted(set(target_map) | set(artifacts)):
            if artifact_id not in target_map:
                content_problems.append(f"{artifact_id}: present but absent from the target")
            elif artifact_id not in artifacts:
                content_problems.append(f"{artifact_id}: missing from the published revision")
            elif artifacts[artifact_id]["blob_hash"] != target_map[artifact_id]:
                content_problems.append(f"{artifact_id}: blob hash differs from the target")
            elif (contents.get(artifact_id) is None or target_contents.get(artifact_id) is None
                  or contents[artifact_id] != target_contents[artifact_id]):
                content_problems.append(f"{artifact_id}: bytes differ or are unreadable")
        if view.get("restored_from") != target_id:
            content_problems.append("published revision does not record restored_from=target")

        changed = sorted(a for a, e in entries.items() if e["changed"])
        removed = sorted(a for a in parent_map if a not in artifacts)
        missing = sorted(a for a in expected_ids
                         if (a in target_map and (a not in artifacts
                                                  or artifacts[a]["blob_hash"] != target_map[a]))
                         or (a not in target_map and a in artifacts))
        unexpected = sorted(a for a in set(artifacts) | set(parent_map)
                            if a not in expected_ids
                            and parent_map.get(a) != (artifacts[a]["blob_hash"]
                                                      if a in artifacts else None))
        expected_edits = [{"artifact_id": a, "edit_kind": "restore",
                           "edits": [{"locator": "$artifact", "before": parent_map.get(a),
                                      "after": target_map.get(a)}]} for a in expected_ids]
        actual_edits = ([{"artifact_id": a, "before_blob_hash": parent_map.get(a),
                          "after_blob_hash": artifacts[a]["blob_hash"],
                          "sha256": entries[a]["sha256"]} for a in changed]
                        + [{"artifact_id": a, "before_blob_hash": parent_map.get(a),
                            "after_blob_hash": None, "sha256": None, "removed": True}
                           for a in removed])
        unexpected_edits = [{"artifact_id": a, "before_blob_hash": parent_map.get(a),
                             "after_blob_hash": (artifacts[a]["blob_hash"]
                                                 if a in artifacts else None),
                             "removed": a not in artifacts} for a in unexpected]
        missing_edits = [{"artifact_id": a,
                          "reason": ("artifact_absent_from_published_revision"
                                     if a not in artifacts and a in target_map
                                     else "artifact_not_removed" if a not in target_map
                                     else "published_blob_differs_from_target")}
                         for a in missing]

        # ---- 3. parse the actual bytes of every restored (changed, present) artifact
        parsed_docs: dict[str, tuple] = {}
        parse_problems, parse_notes = [], []
        for artifact_id in sorted(set(changed) | (set(expected_ids) & set(artifacts))):
            art = artifacts[artifact_id]
            kind = _kind_of(_artifact_name(art), None, False)
            content = contents.get(artifact_id)
            if content is None:
                parse_problems.append(f"{artifact_id}: unreadable or checksum mismatch")
                continue
            if kind is None:
                parse_notes.append(f"{artifact_id}: no parser (not checked)")
                continue
            ok, detail, parsed = _parse(kind, content)
            if ok:
                parsed_docs[artifact_id] = (kind, parsed)
            else:
                parse_problems.append(f"{artifact_id}: {detail}")

        # ---- 4. independent acceptance oracle
        oracles, oracle_meta, oracle_problem = _load_oracles()
        relevant = [o for o in (oracles or []) if o["artifact_id"] in expected_ids]
        oracle_rows = []
        if oracle_problem:
            acceptance = ("NOT_RUN", f"{oracle_problem}: the independent oracle did not run")
        elif not relevant:
            acceptance = ("NOT_RUN", "no independent oracle covers a restored artifact "
                          "(configure RIPPLE_VERIFICATION_ORACLE_FILE)")
        else:
            failures = []
            for oracle in relevant:
                doc = parsed_docs.get(oracle["artifact_id"])
                actual = _MISSING
                if doc is not None and doc[0] == "yaml":
                    actual = _yaml_lookup(doc[1], oracle["path"])
                passed = actual is not _MISSING and _strict_eq(actual, oracle["expected"])
                oracle_rows.append({
                    "oracle_id": oracle["oracle_id"], "artifact_id": oracle["artifact_id"],
                    "path": oracle["path"], "expected": _json_safe(oracle["expected"]),
                    "actual": None if actual is _MISSING else _json_safe(actual),
                    "requirement_ref": oracle.get("requirement_ref"),
                    "status": "PASS" if passed else "FAIL"})
                if not passed:
                    failures.append(oracle["oracle_id"])
            acceptance = (("FAIL", f"oracle(s) disagree: {failures}") if failures
                          else ("PASS", f"{len(relevant)} independent oracle(s) agree"))

        scope = {"target_artifact_ids": expected_ids, "before_blob_hash": None,
                 "after_blob_hash": None, "published_revision_id": published,
                 "verification_plan": [], "review_job_id": None, "review_board_id": None,
                 "review_checks": {}, "restored_from": target_id,
                 "parent_revision_id": parent_id}

        def _check(status, detail, required=True):
            return {"status": status, "required": required, "detail": detail}

        checks = {
            "published_bytes_readback": _check(
                "FAIL" if unreadable else "PASS",
                (f"unreadable: {unreadable}" if unreadable
                 else f"{len(entries)} artifact(s) read back and checksum-verified")),
            "expected_edits": _check(
                "FAIL" if content_problems else "PASS",
                "; ".join(content_problems)
                or "published revision equals the target revision (hashes and bytes)"),
            "parse_after": _check(
                "FAIL" if parse_problems else "PASS",
                "; ".join(parse_problems + parse_notes) or "actual bytes parse"),
            "acceptance_tests": _check(acceptance[0], acceptance[1]),
            "simulation_exact_after": _check(
                "SKIPPED_NOT_APPLICABLE", "a restore republishes existing snapshot bytes",
                required=False),
        }
        reasons = ([f"missing_write:{a}" for a in missing]
                   + [f"unexpected_write:{a}" for a in unexpected]
                   + [f"unreadable_artifact:{a}" for a in unreadable])
        failed = bool(missing or unexpected or unreadable)
        not_run = False
        for name, chk in checks.items():
            if chk["required"] and chk["status"] == "FAIL":
                failed = True
                reasons.append(f"required_check_failed:{name}")
            elif chk["required"] and chk["status"] == "NOT_RUN":
                not_run = True
                reasons.append(f"required_check_not_run:{name}")
        if oracle_problem:
            reasons.append(oracle_problem)
        status = "FAILED" if failed else "NOT_RUN" if not_run else "PASSED"

        body = {
            "id": ctx.job["job_id"],
            "schema_version": REPORT_SCHEMA_VERSION,
            "subject": "RESTORE",
            "apply_job_id": apply_job_id,
            "approval_job_id": None,
            "restore_approval_id": approval["approval_id"],
            "restored_from": target_id,
            "parent_revision_id": parent_id,
            "expected_edits": expected_edits,
            "actual_edits": actual_edits,
            "unexpected_edits": unexpected_edits,
            "missing_edits": missing_edits,
            "parsing_result": {"status": checks["parse_after"]["status"],
                               "problems": parse_problems, "notes": parse_notes},
            "tests_result": {"status": checks["acceptance_tests"]["status"],
                             "oracles": oracle_rows, "oracle_config": oracle_meta},
            "oracle_config": oracle_meta,
            "simulation_result": {"status": checks["simulation_exact_after"]["status"],
                                  "detail": checks["simulation_exact_after"]["detail"]},
            "evidence_result": {"status": checks["published_bytes_readback"]["status"],
                                "artifacts_checked": len(entries), "unreadable": unreadable},
            "open_required_work": [],
            "scope": _json_safe(scope),
            "verification_scope": _json_safe(scope),
            "coverage": {"artifacts_total": len(entries),
                         "artifacts_readable": len(entries) - len(unreadable),
                         "changed": changed, "removed": removed, "expected": expected_ids},
            "after_source_revision_id": published,
            "after_knowledge_revision_id": None,
            "status": status,
            "reasons": reasons,
            "checks": checks,
            "artifacts": {a: entries[a] for a in sorted(entries)},
            "actual_values": {},
            "manual_work": [],
            "scope_mismatches": _scope_mismatch_section(ctx),
        }
        body = _json_safe(body)
        return ExecutionOutcome(result={**body, "report_digest": canonical_hash(body)},
                                model_id=None)


# ----------------------------------------------------------------------------- rebuild

def _completed_jobs(store, tenant_id, workspace_id, kind):
    with store._tx() as cur:
        cur.execute("""SELECT job_id, payload, execution_result, pinned_source_revision_id
            FROM v2_jobs WHERE tenant_id=%s AND workspace_id=%s AND kind=%s
              AND execution_status='COMPLETED' ORDER BY job_id""",
                    (tenant_id, workspace_id, kind))
        rows = cur.fetchall()
    out = []
    for job_id, payload, result, pinned in rows:
        payload = json.loads(payload) if isinstance(payload, str) else (payload or {})
        result = json.loads(result) if isinstance(result, str) else (result or {})
        out.append({"job_id": job_id, "payload": payload, "result": result,
                    "pinned_source_revision_id": pinned})
    return out


def _candidate_ids_by_artifact(extractions: list[dict]) -> dict[str, list[str]]:
    by_artifact: dict[str, list[str]] = {}
    for job in extractions:
        artifact_id = job["payload"].get("artifact_id")
        for cand in (job["result"].get("candidates") or []):
            cid = cand.get("candidate_id") if isinstance(cand, dict) else None
            if artifact_id and cid:
                by_artifact.setdefault(artifact_id, []).append(cid)
    return by_artifact


def _zero_candidate_artifacts(extractions: list[dict]) -> set[str]:
    """Artifacts that HAVE a COMPLETED extraction and whose completed extractions produced no
    candidate at all. Such an artifact has nothing that a relation could depend on, so re-extracting
    its unchanged bytes on every change adds cost and no evidence. An artifact with no completed
    extraction (never run, failed, cancelled) is NOT in this set: `extractions` holds only
    COMPLETED jobs, so an incomplete run stays uncovered."""
    seen: set[str] = set()
    produced: set[str] = set()
    for job in extractions:
        artifact_id = job["payload"].get("artifact_id")
        if not artifact_id:
            continue
        seen.add(artifact_id)
        if any(isinstance(c, dict) and c.get("candidate_id")
               for c in (job["result"].get("candidates") or [])):
            produced.add(artifact_id)
    return seen - produced


def _ensure_stage_job(store, tenant_id, workspace_id, job_id, kind, payload, *, reserve=False):
    try:
        return ensure_job(store, tenant_id, workspace_id, job_id, kind, payload, reserve=reserve)
    except StageIdConflict as exc:
        raise NonRetryableExecutionError("stage_job_id_conflict", str(exc)) from exc


class KmRebuildExecutor:
    """See module docstring. Zero provider calls (`external = False`)."""

    external = False
    executor_id = "km_rebuild.reingest_index_reanalysis.v1"

    def prepare(self, ctx: ExecutionContext) -> ExecutionRequest:
        _payload_ids(ctx)
        return ExecutionRequest(prompt=None,
                                schema={"type": "object", "executor": self.executor_id},
                                provider_calls_planned=0, external=False)

    def _superseded(self, apply_job_id, published, head) -> ExecutionOutcome:
        return ExecutionOutcome(result={
            "executor": self.executor_id, "status": "SUPERSEDED",
            "apply_job_id": apply_job_id, "published_revision_id": published,
            "superseded_by_source_revision_id": head.source_revision_id,
            "reanalysis": {"scope": None, "reason": "superseded_by_newer_source",
                           "artifact_ids": [], "job_ids": []},
            "invalidation": [], "carry_forward": {"carried": [], "reopened": []},
            "counterevidence": [], "knowledge_candidate_revision_id": None,
        }, model_id=None)

    def execute(self, ctx: ExecutionContext, request: ExecutionRequest,
                calls: ProviderCallBudget) -> ExecutionOutcome:
        store, job = ctx.store, ctx.job
        tenant_id, workspace_id = job["tenant_id"], job["workspace_id"]
        apply_job_id, published = _payload_ids(ctx)

        head = store.head(tenant_id, workspace_id)
        if head.source_revision_id != published:
            return self._superseded(apply_job_id, published, head)

        view = _revision_view(ctx, published)
        parent_id = view["parent_revision_id"]
        parent_map = {}
        if parent_id:
            parent_map = {a["artifact_id"]: a["blob_hash"]
                          for a in _revision_view(ctx, parent_id)["artifacts"]}
        artifacts = view["artifacts"]
        changed = sorted(a["artifact_id"] for a in artifacts
                         if parent_map.get(a["artifact_id"]) != a["blob_hash"])
        # An artifact a RESTORE removed no longer exists in the new
        # revision; its candidates are as affected as a changed one's.
        removed_ids = sorted(a for a in parent_map
                             if a not in {x["artifact_id"] for x in artifacts})

        # ---- 1. NEW observations under the NEW revision id (old IR untouched; idempotent)
        observation_counts: dict[str, int] = {}
        coverage_complete = True
        for art in artifacts:
            try:
                content = store.read_blob(tenant_id, workspace_id, published, art["artifact_id"])
            except (BlobIntegrityError, KeyError, FileNotFoundError, TenantScopeError) as exc:
                raise NonRetryableExecutionError(
                    "source_blob_integrity_failed", art["artifact_id"]) from exc
            name = _artifact_name(art)
            parser = parser_for(name)
            if parser is None:
                coverage_complete = False
                observation_counts[art["artifact_id"]] = 0
                continue
            parsed = parser(name, content)
            if parsed.parse_status != "parsed":
                coverage_complete = False
            observations = list(parsed.observations)
            if not any(o.get("kind") == "text_chunk" for o in observations):
                # A structured adapter (yaml/csv/jsonl) records values, not the
                # source text. Keep the readable source text as ONE chunk (the
                # shape `text_chunk` observations already have) so retrieval
                # and the citation resolver can find the NEW text by key.
                try:
                    text = content.decode("utf-8-sig")
                except UnicodeDecodeError:
                    text = None
                if text and text.strip():
                    observations.append({
                        "kind": "text_chunk",
                        "locator": {"type": "line_span", "line_start": 1,
                                    "line_end": len(text.splitlines())},
                        "raw_value": text, "fields": {}, "parser_version": PARSER_VERSION})
            if observations:
                store.record_source_observations(
                    tenant_id, workspace_id, published, art["artifact_id"], observations)
            observation_counts[art["artifact_id"]] = len(observations)

        # ---- 2. retrieval index for the NEW revision (failure => named FAILED, retryable)
        try:
            retrieval_index = store.build_retrieval_index(
                tenant_id, workspace_id, published,
                source_coverage_complete=coverage_complete,
                source_coverage_reason=("all artifacts of the published revision parsed"
                                        if coverage_complete else
                                        "at least one artifact of the published revision "
                                        "has no parser or did not fully parse"))
        except Exception as exc:  # noqa: BLE001 - deterministic index failure => terminal FAILED
            raise NonRetryableExecutionError(
                "index_build_failed", f"{type(exc).__name__}: {exc}") from exc

        # ---- 3. a newer source landed while re-ingesting: nothing more is scheduled
        head = store.head(tenant_id, workspace_id)
        if head.source_revision_id != published:
            return self._superseded(apply_job_id, published, head)

        extractions = _completed_jobs(store, tenant_id, workspace_id, "extraction")
        relations = _completed_jobs(store, tenant_id, workspace_id, "relation")
        by_artifact = _candidate_ids_by_artifact(extractions)

        # ---- 4. affected claims: candidates of the changed artifacts (older revisions)
        #         go through the EXISTING invalidation executor.
        decision_id = apply_decision_id(apply_job_id)
        stale_candidates = sorted({
            cid for job_ in extractions
            if job_["payload"].get("artifact_id") in (changed + removed_ids)
            and job_["pinned_source_revision_id"] != published
            for cid in _candidate_ids_by_artifact([job_]).get(job_["payload"]["artifact_id"], [])})
        invalidation = []
        for cid in stale_candidates:
            inv_job_id = f"inval-{apply_job_id}-{cid}"
            _ensure_stage_job(store, tenant_id, workspace_id, inv_job_id, "invalidation",
                       {"decision_id": decision_id, "target": cid})
            dependents = store.find_dependent_jobs_for_candidate(tenant_id, workspace_id, cid)
            invalidation.append({"candidate_id": cid, "job_id": inv_job_id,
                                 "dependents": [{"job_id": d["job_id"], "kind": d["kind"]}
                                                for d in dependents]})

        # ---- 5. reanalysis scope: FULL when dependency coverage is missing
        covered = set()
        for rel in relations:
            covered |= set(rel["payload"].get("candidate_ids") or [])
        zero_candidate = _zero_candidate_artifacts(extractions)
        uncovered = sorted(a["artifact_id"] for a in artifacts
                           if a["artifact_id"] not in changed
                           and a["artifact_id"] not in zero_candidate
                           and not (set(by_artifact.get(a["artifact_id"], [])) & covered))
        if uncovered:
            scope = "FULL"
            reason = f"dependency_coverage_missing: {uncovered}"
            # Same set as before; only the ORDER is chosen: the changed artifacts are re-read first, so a project cost cap that stops the wave
            # early leaves the unchanged files (not the changed ones) for later.
            targets = sorted((a["artifact_id"] for a in artifacts), key=lambda x: (x not in changed, x))
        else:
            scope = "TARGETED"
            reason = "dependency_coverage_available_for_all_unchanged_artifacts"
            targets = changed
        reanalysis_job_ids, deferred_for_budget = [], []
        for i, artifact_id in enumerate(targets):
            rj = f"reanalysis-apply-{apply_job_id}-{artifact_id}"
            try:
                _ensure_stage_job(store, tenant_id, workspace_id, rj, "extraction",
                                  {"artifact_id": artifact_id, "caused_by_apply_job_id": apply_job_id}, reserve=True)
            except BudgetExceeded:
                # The re-extraction is NOT cut to fit and knowledge is NOT called fresh: the artifacts that could not be reserved are reported
                # (loop status shows them as unfinished) and stay unanalysed on the new revision until the cap allows them.
                deferred_for_budget = list(targets[i:])
                break
            reanalysis_job_ids.append(rj)

        # ---- 6. decision carry-forward: only when the evidence binding validates
        carried, reopened = [], []
        with store._tx() as cur:
            cur.execute("""SELECT decision_id, target, action, evidence_ref, provenance
                FROM v2_decisions WHERE tenant_id=%s AND workspace_id=%s
                  AND superseded_by IS NULL AND revoked_by IS NULL ORDER BY epoch""",
                        (tenant_id, workspace_id))
            decisions = cur.fetchall()
        for d_id, target_, action, evidence_ref, provenance in decisions:
            row = {"decision_id": d_id, "target": target_, "action": action}
            ref = None
            if evidence_ref:
                try:
                    ref = json.loads(evidence_ref)
                except ValueError:
                    ref = None
            if not isinstance(ref, dict):
                reopened.append({**row, "reason": "no_evidence_binding: "
                                 f"provenance={provenance or 'human_assertion'}"})
                continue
            try:
                report = store.validate_evidence(tenant_id, workspace_id, published, [ref])[0]
            except Exception as exc:  # noqa: BLE001 - an unvalidatable binding is never carried
                reopened.append({**row, "reason": f"evidence_binding_unvalidatable: {exc}"})
                continue
            if report.get("citation_valid"):
                carried.append({**row, "evidence_binding": "valid_on_new_source"})
            else:
                reopened.append({**row, "reason": "evidence_binding_invalid: "
                                 f"{report.get('reason')}"})

        # ---- 7. counterevidence: NEW index only, content that still asserts the OLD value
        counterevidence = []
        seen = set()
        proposal_edits = _proposal_edits(ctx, apply_job_id)
        for edit in proposal_edits:
            before = edit.get("before")
            if before is None:
                continue
            leaf = str(edit.get("path") or edit.get("cell") or "").split(".")[-1]
            found = store.retrieve(tenant_id, workspace_id, f"{leaf} {before}", top_k=10)
            if found.get("source_revision_id") != published or found.get("status") != "READY":
                continue
            for hit in found.get("results", []):
                if str(before) not in json.dumps(hit.get("raw_value"), default=str):
                    continue
                key = (hit["artifact_id"], hit["observation_id"])
                if key in seen:
                    continue
                seen.add(key)
                counterevidence.append(_json_safe({
                    "artifact_id": hit["artifact_id"], "observation_id": hit["observation_id"],
                    "blob_hash": hit.get("blob_hash"), "raw_value": hit.get("raw_value"),
                    "score": hit.get("score"), "old_value": before}))

        result = {
            "executor": self.executor_id, "status": "COMPLETED",
            "apply_job_id": apply_job_id, "published_revision_id": published,
            "parent_revision_id": parent_id, "changed_artifact_ids": changed,
            "observations": observation_counts,
            "retrieval_index": _json_safe(retrieval_index),
            "reanalysis": {"scope": scope, "reason": reason, "artifact_ids": targets,
                           "job_ids": reanalysis_job_ids, "budget_deferred_artifact_ids": deferred_for_budget,
                           "complete": not deferred_for_budget},
            "invalidation": invalidation,
            "carry_forward": {"carried": carried, "reopened": reopened},
            "counterevidence": counterevidence,
            "knowledge_candidate_revision_id": knowledge_candidate_id(apply_job_id),
        }
        return ExecutionOutcome(result=_json_safe(result), model_id=None)


def _proposal_edits(ctx, apply_job_id: str) -> list[dict]:
    """The approved proposal's recorded typed edits (read-only; used only to
    phrase the counterevidence query)."""
    store, tenant_id, workspace_id = ctx.store, ctx.job["tenant_id"], ctx.job["workspace_id"]
    try:
        apply_job = store.job_record(tenant_id, workspace_id, apply_job_id)
        approval = store.job_record(
            tenant_id, workspace_id, (apply_job.get("payload") or {}).get("approval_job_id"))
        proposal = store.job_record(
            tenant_id, workspace_id, (approval.get("execution_result") or {}).get("proposal_job_id"))
    except TenantScopeError:
        return []
    return [dict(e) for e in ((proposal.get("execution_result") or {}).get("edits") or [])
            if isinstance(e, dict)]


__all__ = ["VerificationExecutor", "KmRebuildExecutor", "consume_source_published",
           "ensure_job", "verification_job_id", "rebuild_job_id", "apply_decision_id",
           "knowledge_candidate_id"]
