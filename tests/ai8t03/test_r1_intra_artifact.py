"""Intra-artifact verification: an UNEXPECTED WRITE INSIDE the target artifact.

The apply gate published the artifact the plan named, but the durable bytes carry
more than (parent bytes + approved edits): an extra key, a changed unrelated key,
a removed key, an extra / changed xlsx cell, an extra sheet, an extra / changed
defined name, or a bool written where the plan says an int. The verifier must
compare the FULL published content and report FAILED with `unexpected_edits`
(scope `content`, hashes, per-locator differences); a legitimate edit and a
serialization-only re-save must still be accepted (no false positive).

Every scenario goes through the real chain (apply job -> publish ->
SOURCE_PUBLISHED -> dispatcher -> verification job). The fault is injected at the
WRITE / STORAGE boundary only (a tampered `publish_source` content, a deleted
parent blob, a mutated recorded edit), never by editing the verifier. Each
negative has a positive control in which only the approved edit differs.
"""

from __future__ import annotations

import io
from uuid import uuid4

import pytest

from tests.ai8t03.conftest import (
    applied_change, change_dto, install_provider, mutate_job_result, pump, reasons_of,
    report_of, write_oracle,
)

NOT_DONE = ("MACHINE_VERIFIED", "AWAITING_ACCEPTANCE", "COMPLETED")
BASE = "refund_window_days: 14\nkeep_a: 1\nkeep_b: two\n"


def _target():
    return f"config-r1-{uuid4().hex[:6]}.yaml"


def _read(store, content, key):
    return store.blobs.read(content[key])


def _yaml_change(setup, monkeypatch, owner, tamper=None, *, siblings=None, proposed=30):
    return applied_change(setup, monkeypatch, owner, target=_target(), siblings=siblings,
                          publish_tamper=tamper, base_text=BASE, proposed=proposed)


def _finish(setup, monkeypatch, tmp_path, change, *, oracle=30):
    if oracle is not None:
        write_oracle(tmp_path, monkeypatch, artifact_id=change["target"], expected=oracle)
    install_provider(monkeypatch, "ok")
    pump(setup)
    return report_of(change, setup["store"], setup["tenant"])


def _content_diffs(report, artifact_id):
    entry = [e for e in report["unexpected_edits"] if e["artifact_id"] == artifact_id]
    assert entry, ("no unexpected_edits entry", report["unexpected_edits"])
    return entry[0]


def _assert_gate_blocked(setup, change, report):
    assert report["status"] == "FAILED", (report["status"], report["reasons"])
    dto = change_dto(change["client"], change["workspace"], change["apply_job_id"])
    assert dto["state"] not in NOT_DONE, dto
    assert dto["completion"]["eligible"] is False
    resp = change["ready"]["approver"]["client"].post(
        f"/v2/changes/{change['apply_job_id']}/result-decision", json={
            "workspace_id": change["workspace"], "action": "accept", "reason": "attempt",
            "expected_report_digest": report["report_digest"]})
    assert resp.status_code == 409 and "machine_verification_not_pass" in resp.text, resp.text


# ------------------------------------------------------------------ yaml intra-artifact

def _replace(text, old, new):
    assert old in text, (old, text)
    return text.replace(old, new)


YAML_TAMPERS = {
    "extra-key": (lambda t: t + "extra_key: 1\n", "added", "extra_key"),
    "changed-unrelated-key": (lambda t: _replace(t, "keep_a: 1", "keep_a: 2"), "changed", "keep_a"),
    "removed-key": (lambda t: _replace(t, "keep_b: two\n", ""), "removed", "keep_b"),
}


@pytest.mark.parametrize("variant", sorted(YAML_TAMPERS))
def test_r1_yaml_unexpected_write_inside_the_target_fails_with_scope_and_hashes(
        setup, monkeypatch, owner, tmp_path, variant):
    mutate, kind, locator = YAML_TAMPERS[variant]

    def tamper(store, content):
        key = next(iter(content))
        text = _read(store, content, key).decode("utf-8")
        return {**content, key: store.stage_blob(mutate(text).encode("utf-8"))}

    change = _yaml_change(setup, monkeypatch, owner, tamper)
    report = _finish(setup, monkeypatch, tmp_path, change)
    target = change["target"]
    entry = _content_diffs(report, target)
    assert entry["scope"] == "content"
    assert entry["before_blob_hash"] and entry["after_blob_hash"] and entry["sha256"]
    assert {"locator": locator, "kind": kind} == {
        k: entry["differences"][0][k] for k in ("locator", "kind")} or any(
        d["locator"] == locator and d["kind"] == kind for d in entry["differences"]), entry
    assert f"unexpected_write:{target}" in reasons_of(report)
    assert report["missing_edits"] == []  # the approved edit itself DID land
    _assert_gate_blocked(setup, change, report)
    # the body (and therefore the digest) carries the finding: digest is stable and non-empty
    assert report["report_digest"]


def test_r1_yaml_positive_control_only_the_edited_key_differs_passes(
        setup, monkeypatch, owner, tmp_path):
    change = _yaml_change(setup, monkeypatch, owner)
    report = _finish(setup, monkeypatch, tmp_path, change)
    assert report["status"] == "PASSED", report["reasons"]
    assert report["unexpected_edits"] == [] and report["missing_edits"] == []
    dto = change_dto(change["client"], change["workspace"], change["apply_job_id"])
    assert dto["verification"]["status"] == "PASSED"


def test_r1_yaml_key_order_only_change_is_not_flagged(setup, monkeypatch, owner, tmp_path):
    """Serialization-only difference (key order): the structural comparison must
    not report an unexpected write."""
    def tamper(store, content):
        key = next(iter(content))
        return {**content, key: store.stage_blob(
            b"keep_b: two\nkeep_a: 1\nrefund_window_days: 30\n")}

    change = _yaml_change(setup, monkeypatch, owner, tamper)
    report = _finish(setup, monkeypatch, tmp_path, change)
    assert report["unexpected_edits"] == [], report["unexpected_edits"]
    assert report["status"] == "PASSED", report["reasons"]


def test_r1_bool_written_for_an_int_is_a_failure_and_the_int_control_passes(
        setup, monkeypatch, owner, tmp_path):
    """R-2 (True vs 1): the plan proposes 1; the bytes say `true`."""
    def tamper(store, content):
        key = next(iter(content))
        return {**content, key: store.stage_blob(
            b"refund_window_days: true\nkeep_a: 1\nkeep_b: two\n")}

    change = _yaml_change(setup, monkeypatch, owner, tamper, proposed=1)
    report = _finish(setup, monkeypatch, tmp_path, change, oracle=1)
    assert report["status"] == "FAILED", report["reasons"]
    assert report["checks"]["expected_edits"]["status"] == "FAIL"
    assert report["checks"]["acceptance_tests"]["status"] == "FAIL"
    _assert_gate_blocked(setup, change, report)


def test_r1_int_control_for_the_bool_case_passes(setup, monkeypatch, owner, tmp_path):
    change = _yaml_change(setup, monkeypatch, owner, proposed=1)
    report = _finish(setup, monkeypatch, tmp_path, change, oracle=1)
    assert report["status"] == "PASSED", report["reasons"]


# ------------------------------------------------------------------ multi-artifact / text

def test_r1_multi_artifact_head_intra_target_write_is_scoped_to_the_target_only(
        setup, monkeypatch, owner, tmp_path):
    siblings = {"sibling-keep.yaml": "keep: 1\n", "notes.txt": "plain text sibling\n"}

    def tamper(store, content):
        key = next(k for k in content if k.startswith("config-r1-"))
        text = _read(store, content, key).decode("utf-8") + "extra_key: 1\n"
        return {**content, key: store.stage_blob(text.encode("utf-8"))}

    change = _yaml_change(setup, monkeypatch, owner, tamper, siblings=siblings)
    report = _finish(setup, monkeypatch, tmp_path, change)
    assert [e["artifact_id"] for e in report["unexpected_edits"]] == [change["target"]]
    assert report["artifacts"]["sibling-keep.yaml"]["checksum_ok"] is True
    _assert_gate_blocked(setup, change, report)


def test_r1_multi_artifact_positive_control_untouched_siblings_pass(
        setup, monkeypatch, owner, tmp_path):
    siblings = {"sibling-keep.yaml": "keep: 1\n", "notes.txt": "plain text sibling\n"}
    change = _yaml_change(setup, monkeypatch, owner, siblings=siblings)
    report = _finish(setup, monkeypatch, tmp_path, change)
    assert report["status"] == "PASSED" and report["unexpected_edits"] == [], report["reasons"]


def test_r1_text_artifact_content_change_is_an_unexpected_write(
        setup, monkeypatch, owner, tmp_path):
    def tamper(store, content):
        return {**content, "notes.txt": store.stage_blob(b"plain text sibling, edited\n")}

    change = _yaml_change(setup, monkeypatch, owner, tamper,
                          siblings={"notes.txt": "plain text sibling\n"})
    report = _finish(setup, monkeypatch, tmp_path, change)
    assert "notes.txt" in [e["artifact_id"] for e in report["unexpected_edits"]]
    assert "unexpected_write:notes.txt" in reasons_of(report)
    _assert_gate_blocked(setup, change, report)


# ------------------------------------------------------------------ cannot build expected state

def test_r1_unreadable_parent_bytes_give_a_difference_never_passed(
        setup, monkeypatch, owner, tmp_path):
    change = _yaml_change(setup, monkeypatch, owner)
    store, tenant, ws = setup["store"], setup["tenant"], setup["workspace"]
    parent = {a["artifact_id"]: a["blob_hash"] for a in
              store.source_revision_view(tenant, ws, change["s0"])["artifacts"]}[change["target"]]
    snapshot = dict(store.blobs.data)
    try:
        store.blobs.data.pop(parent)
        report = _finish(setup, monkeypatch, tmp_path, change)
    finally:
        store.blobs.data.clear()
        store.blobs.data.update(snapshot)
    assert report["status"] != "PASSED", report["reasons"]
    dto = change_dto(change["client"], ws, change["apply_job_id"])
    assert dto["state"] not in NOT_DONE and dto["completion"]["eligible"] is False
    joined = str(report["unexpected_edits"]) + reasons_of(report)
    assert "expected_state_unbuildable" in joined or "unexpected_write" in joined, report


def test_r1_stale_recorded_edit_gives_a_difference_never_passed(
        setup, monkeypatch, owner, tmp_path):
    change = _yaml_change(setup, monkeypatch, owner)
    store, tenant, ws = setup["store"], setup["tenant"], setup["workspace"]

    def stale(result):
        for e in result["edits"]:
            e["before"] = 99  # the recorded 'before' no longer matches the parent bytes

    mutate_job_result(store, tenant, ws, change["ready"]["proposal_job_id"], stale)
    report = _finish(setup, monkeypatch, tmp_path, change)
    assert report["status"] != "PASSED", report["reasons"]
    joined = str(report["unexpected_edits"]) + str(report["reasons"])
    assert "expected_state_unbuildable" in joined or report["checks"]["expected_edits"][
        "status"] == "FAIL", report
    dto = change_dto(change["client"], ws, change["apply_job_id"])
    assert dto["completion"]["eligible"] is False


# ------------------------------------------------------------------ xlsx

def _named_workbook_bytes(with_name=True):
    import openpyxl
    from openpyxl.workbook.defined_name import DefinedName
    from tests.ai8t01.conftest import make_xlsx_bytes
    wb = openpyxl.load_workbook(io.BytesIO(make_xlsx_bytes()))
    if with_name:
        wb.defined_names["RateCell"] = DefinedName("RateCell", attr_text="Sheet1!$A$1")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xlsx_change(setup, tamper=None, *, with_name=True):
    from tests.ai8t01.conftest import (
        build_real_ready_xlsx_approval, record_real_approve_decision,
        submit_claim_and_get_apply,
    )
    store, tenant, ws = setup["store"], setup["tenant"], setup["workspace"]
    client = setup["browser"]()
    from tests.ai4t01.conftest import login
    login(client, setup["requester"], setup["credential"])
    target = f"rates-r1-{uuid4().hex[:6]}.xlsx"
    approval_job_id, body, proposal_job_id, review_job_id, _orig = build_real_ready_xlsx_approval(
        setup, client, ws, target_artifact_id=target, workbook_bytes=_named_workbook_bytes(with_name))
    digest = body["execution_result"]["approval_digest"]
    approver = record_real_approve_decision(setup, ws, approval_job_id=approval_job_id,
                                            digest=digest)
    before = store.head(tenant, ws)
    apply_job_id = f"job-apply-{uuid4().hex}"
    with monkeypatch_ctx() as mp:
        if tamper is not None:
            from infra.adapters import v2_postgres
            original = v2_postgres.PostgresV2Store.publish_source

            def _tampered(self, tenant_id, workspace_id, revision_id, content, *a, **k):
                if k.get("operation") == "apply.publish":
                    content = tamper(self, dict(content))
                return original(self, tenant_id, workspace_id, revision_id, content, *a, **k)
            mp.setattr(v2_postgres.PostgresV2Store, "publish_source", _tampered)
        applied = submit_claim_and_get_apply(
            setup, client, ws, approval_job_id=approval_job_id, job_id=apply_job_id,
            expected_approval_digest=digest)
    assert applied["execution_status"] == "COMPLETED", applied
    after = store.head(tenant, ws)
    return {"client": client, "workspace": ws, "apply_job_id": apply_job_id,
            "verify_job_id": f"verify-{apply_job_id}", "target": target,
            "s0": before.source_revision_id, "s1": after.source_revision_id,
            "ready": {"approver": approver, "proposal_job_id": proposal_job_id}}


def monkeypatch_ctx():
    return pytest.MonkeyPatch.context()


def _edit_workbook(store, content, fn):
    import openpyxl
    key = next(iter(content))
    wb = openpyxl.load_workbook(io.BytesIO(store.blobs.read(content[key])))
    fn(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return {**content, key: store.stage_blob(buf.getvalue())}


def _defined_change(wb):
    from openpyxl.workbook.defined_name import DefinedName
    wb.defined_names["RateCell"] = DefinedName("RateCell", attr_text="Sheet1!$A$2")


def _defined_extra(wb):
    from openpyxl.workbook.defined_name import DefinedName
    wb.defined_names["RogueName"] = DefinedName("RogueName", attr_text="Notes!$A$1")


XLSX_TAMPERS = {
    "extra-cell": (lambda wb: wb["Notes"].__setitem__("C5", "rogue"), "Notes!C5"),
    "changed-cell": (lambda wb: wb["Notes"].__setitem__("A1", 99), "Notes!A1"),
    "extra-sheet": (lambda wb: wb.create_sheet("Rogue"), "$sheets"),
    "extra-defined-name": (_defined_extra, "$name:RogueName"),
    "changed-defined-name": (_defined_change, "$name:RateCell"),
}


@pytest.mark.parametrize("variant", sorted(XLSX_TAMPERS))
def test_r1_xlsx_unexpected_write_inside_the_target_fails(
        setup, monkeypatch, owner, tmp_path, variant):
    fn, locator = XLSX_TAMPERS[variant]
    change = _xlsx_change(setup, lambda store, content: _edit_workbook(store, content, fn))
    install_provider(monkeypatch, "ok")
    pump(setup)
    report = report_of(change, setup["store"], setup["tenant"])
    entry = _content_diffs(report, change["target"])
    assert entry["scope"] == "content" and entry["sha256"]
    assert any(d["locator"] == locator for d in entry["differences"]), entry["differences"]
    assert report["status"] == "FAILED", report["reasons"]
    assert f"unexpected_write:{change['target']}" in reasons_of(report)
    dto = change_dto(change["client"], change["workspace"], change["apply_job_id"])
    assert dto["state"] not in NOT_DONE and dto["completion"]["eligible"] is False


def test_r1_xlsx_positive_control_and_normalization_only_resave_are_not_flagged(
        setup, monkeypatch, owner, tmp_path):
    # (a) clean approved edit only
    change = _xlsx_change(setup)
    install_provider(monkeypatch, "ok")
    pump(setup)
    report = report_of(change, setup["store"], setup["tenant"])
    assert report["unexpected_edits"] == [], report["unexpected_edits"]
    assert report["checks"]["expected_edits"]["status"] == "PASS", report["checks"]
    assert report["status"] != "FAILED", report["reasons"]
    # (b) a normalization-only re-save (load + save, no semantic change)
    change2 = _xlsx_change(setup, lambda store, content: _edit_workbook(
        store, content, lambda wb: None))
    pump(setup)
    report2 = report_of(change2, setup["store"], setup["tenant"])
    assert report2["unexpected_edits"] == [], report2["unexpected_edits"]
    assert report2["checks"]["expected_edits"]["status"] == "PASS", report2["checks"]
    assert report2["status"] != "FAILED", report2["reasons"]
