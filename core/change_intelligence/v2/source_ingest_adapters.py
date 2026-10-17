"""Pure, byte-in adapters that turn one admitted artifact's bytes
into typed observations plus a coverage-manifest entry.

Every adapter here is a pure function of `(logical_path, content: bytes)`.
No filesystem, network, subprocess or database access; no macro/formula/
Python execution (contracts C09, support matrix §1). Formula text and any
cached value in a workbook are kept as separate fields with an explicit
recalc status; a stale cached value is never reported as current. A record
(CSV/JSONL row) is an observation, never promoted to a rule.

Reuse decision: `core/adapters/*`
and `core/change_intelligence/ingest.py` target RULE-CANDIDATE identity
resolution (`Candidate`/`IRRecord` with `semantic_hints`), a different
downstream consumer than this module's observation/coverage-manifest contract.
Building dedicated, byte-pure observation adapters here avoids bending that
pipeline's shape (and its `semantic_hints` parameter, which C17 forbids on
this path) to fit a different DTO; `openpyxl`/`yaml`/`ast`/`csv`/`json` are
the same underlying libraries every existing adapter already uses.
"""

from __future__ import annotations

import ast
import csv
import datetime
import hashlib
import io
import json
from dataclasses import dataclass, field

import yaml


class UnrepresentableValueError(ValueError):
    """Raised when a parsed value cannot be turned into a JSON-safe form.
    Caught per-artifact by the calling `parse_*` function so one
    bad value fails that artifact only, with a located diagnostic, rather
    than crashing the worker with an unhandled JSON-serialization error
    downstream at the persistence boundary."""


def _json_safe(value):
    """Canonically encode a parsed value so it is always persistable as
    JSONB. A YAML/XLSX date or datetime is not silently dropped or
    left as a Python object the store cannot serialize -- it is encoded as
    its canonical ISO-8601 string, still readable and still typed by
    convention. A value with no safe encoding (e.g. a Python `complex` or
    `set` literal) raises `UnrepresentableValueError` instead of reaching
    the database layer."""
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, dict):
        return {str(k): _json_safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(v) for v in value]
    raise UnrepresentableValueError(
        f"value of type {type(value).__name__!r} has no JSON-safe encoding")

#: Bumped whenever an adapter's observation shape changes; recorded on every
#: observation row and every coverage-manifest artifact entry.
PARSER_VERSION = "1.0.0"


@dataclass
class ParseResult:
    parse_status: str  # "parsed" | "partial" | "failed"
    parser: str | None
    encoding: str | None
    newline: str | None
    observations: list[dict]
    diagnostics: list[dict]
    #: Adapters keep PARSER_VERSION unless their observation shape grew (Python: see PYTHON_PARSER_VERSION).
    parser_version: str = PARSER_VERSION

    def manifest_entry(self, artifact_id: str, logical_path: str) -> dict:
        return {
            "artifact_id": artifact_id, "logical_path": logical_path,
            "parse_status": self.parse_status, "parser": self.parser,
            "parser_version": self.parser_version, "encoding": self.encoding,
            "newline": self.newline, "diagnostics": self.diagnostics,
            "observation_count": len(self.observations),
            "chunk_count": sum(1 for o in self.observations if o["kind"] == "text_chunk"),
        }


def _diag(code: str, message: str, **locator) -> dict:
    return {"code": code, "message": message, "locator": locator}


def _obs(kind: str, locator: dict, raw_value, fields: dict | None = None, *, version: str = PARSER_VERSION) -> dict:
    return {"kind": kind, "locator": locator, "raw_value": _json_safe(raw_value),
            "fields": _json_safe(dict(fields or {})), "parser_version": version}


def _decode(content: bytes) -> tuple[str, str, str]:
    """Returns (text, encoding, newline). Raises UnicodeDecodeError if the
    bytes are not valid UTF-8 (admission already rejected undecodable text
    for observation-eligible formats, so this is a defensive re-check)."""
    text = content.decode("utf-8-sig")
    newline = "crlf" if "\r\n" in text else "lf"
    return text, "utf-8", newline


def _lines(text: str, newline: str) -> list[str]:
    return text.split("\r\n" if newline == "crlf" else "\n")


# ----------------------------------------------------------------- markdown

def parse_markdown(logical_path: str, content: bytes) -> ParseResult:
    text, encoding, newline = _decode(content)
    sep = "\r\n" if newline == "crlf" else "\n"
    lines = _lines(text, newline)
    observations = []
    start = None
    for i, line in enumerate(lines, start=1):
        if line.strip():
            if start is None:
                start = i
        elif start is not None:
            observations.append(_chunk(lines, start, i - 1, sep))
            start = None
    if start is not None:
        observations.append(_chunk(lines, start, len(lines), sep))
    return ParseResult("parsed", "markdown.chunks.v1", encoding, newline, observations, [])


def _chunk(lines: list[str], start: int, end: int, sep: str) -> dict:
    raw = sep.join(lines[start - 1:end])
    return _obs("text_chunk", {"type": "line_span", "line_start": start, "line_end": end}, raw)


# ----------------------------------------------------------------- yaml

def parse_yaml(logical_path: str, content: bytes) -> ParseResult:
    text, encoding, newline = _decode(content)
    try:
        docs = list(yaml.safe_load_all(text))
    except yaml.YAMLError as exc:
        mark = getattr(exc, "problem_mark", None)
        line = (mark.line + 1) if mark is not None else 1
        return ParseResult("failed", "yaml.values.v1", encoding, newline, [],
                           [_diag("invalid_yaml", str(exc), line_start=line)])
    if len(docs) != 1:
        return ParseResult("failed", "yaml.values.v1", encoding, newline, [],
                           [_diag("multi_document_yaml",
                                  "a single YAML document is required", line_start=1)])
    observations, diagnostics = [], []
    for key_path, value in _yaml_leaves(docs[0] or {}, ()):
        try:
            observations.append(_obs("yaml_value", {"type": "key_path", "key_path": key_path}, value))
        except UnrepresentableValueError as exc:
            diagnostics.append(_diag("unrepresentable_value", str(exc), key_path=key_path))
    status = "failed" if observations == [] and diagnostics else (
        "partial" if diagnostics else "parsed")
    return ParseResult(status, "yaml.values.v1", encoding, newline, observations, diagnostics)


def _yaml_leaves(value, prefix: tuple):
    if isinstance(value, dict):
        for key in sorted(value, key=str):
            yield from _yaml_leaves(value[key], prefix + (str(key),))
        return
    yield ".".join(prefix), value


# ----------------------------------------------------------------- python (AST only, never executed)

#: Python observations carry their OWN parser version: the shape grew (code blocks were added next to the literal assignments, which keep exactly their
#: previous locator and value). Observation ids embed the version, an old revision's rows stay readable, and the store reads the highest version per artifact.
PYTHON_PARSER_VERSION = "1.1.0"
#: Bounds (a block is what a model is shown and what a quote is resolved against, so it is size-limited; an over-long function is cut into consecutive parts).
MAX_PY_SOURCE_BYTES = 256 * 1024
MAX_BLOCK_LINES = 80
MAX_BLOCK_CHARS = 6000
MAX_CODE_BLOCKS = 400


def _node_start(node) -> int:
    return min([node.lineno] + [d.lineno for d in getattr(node, "decorator_list", [])])


def _with_leading_comments(lines: list[str], start: int) -> int:
    """A def/class keeps the comment lines written directly above it (no blank line between): they are part of its original text and often state its intent."""
    while start > 1 and lines[start - 2].lstrip().startswith("#"):
        start -= 1
    return start


def _is_docstring(stmt) -> bool:
    return isinstance(stmt, ast.Expr) and isinstance(getattr(stmt, "value", None), ast.Constant) and isinstance(stmt.value.value, str)


def _is_pure_literal_assign(node) -> bool:
    if not isinstance(node, ast.Assign):
        return False
    try:
        ast.literal_eval(node.value)
    except (ValueError, TypeError, SyntaxError, MemoryError, RecursionError):
        return False
    return all(isinstance(t, ast.Name) for t in node.targets)


def _segments(lines: list[str], sep: str, start: int, end: int, breaks: list[int]) -> list[tuple[int, int]]:
    """Cut [start, end] into consecutive line ranges that respect the size bounds, preferring statement starts (`breaks`) and never splitting inside a line."""
    def fits(a: int, b: int) -> bool:
        return (b - a + 1) <= MAX_BLOCK_LINES and len(sep.join(lines[a - 1:b])) <= MAX_BLOCK_CHARS
    if fits(start, end):
        return [(start, end)]
    cuts = sorted({b for b in breaks if start < b <= end})
    out, a = [], start
    while a <= end:
        if fits(a, end):
            out.append((a, end))
            break
        best = None
        for c in cuts:
            if c <= a:
                continue
            if fits(a, c - 1):
                best = c - 1
            else:
                break
        if best is None:                          # one statement (or one run of lines) larger than the bound: cut by lines
            best = a
            while best < end and fits(a, best + 1):
                best += 1
        out.append((a, best))
        a = best + 1
    return out


def _code_blocks(tree, lines: list[str], sep: str) -> tuple[list[dict], bool]:
    """Observations for the code itself: each module-level function, each method, a class header and the remaining module-level statements, as the ORIGINAL
    source text of an exact line span. No analysis is claimed: no call graph, no types, no cross-file links."""
    body, blocks = tree.body, []
    truncated = [False]

    def emit(kind: str, qualname: str, start: int, end: int, breaks: list[int], doc):
        for i, (a, b) in enumerate(_segments(lines, sep, start, end, breaks), start=1):
            if len(blocks) >= MAX_CODE_BLOCKS:
                truncated[0] = True
                return
            raw = sep.join(lines[a - 1:b])
            loc = {"type": "code_span", "qualname": qualname, "line_start": a, "line_end": b}
            fields = {"node": kind, "qualname": qualname, "docstring": bool(doc) and i == 1,
                      "text_sha256": hashlib.sha256(raw.encode("utf-8")).hexdigest()}
            if doc and i == 1:
                fields["docstring_lines"] = [doc[0], doc[1]]
            parts = _segments(lines, sep, start, end, breaks)
            if len(parts) > 1:
                loc["part"], loc["parts"] = i, len(parts)
                fields["part"], fields["parts"] = i, len(parts)
            blocks.append(_obs("code_block", loc, raw, fields))

    def doc_span(node):
        first = node.body[0] if getattr(node, "body", None) else None
        return (first.lineno, first.end_lineno) if first is not None and _is_docstring(first) else None

    def fn(node, qualname, kind):
        emit(kind, qualname, _with_leading_comments(lines, _node_start(node)), node.end_lineno, [_node_start(s) for s in node.body], doc_span(node))

    run: list = []

    def flush_run():
        if run:
            kind = "module_docstring" if len(run) == 1 and _is_docstring(run[0]) and run[0] is body[0] else "module_code"
            emit(kind, "<module>", run[0].lineno, run[-1].end_lineno, [s.lineno for s in run], None)
            run.clear()

    def class_block(node):
        start = _with_leading_comments(lines, _node_start(node))
        first_method = next((s for s in node.body if isinstance(s, (ast.FunctionDef, ast.AsyncFunctionDef))), None)
        header_end = (_with_leading_comments(lines, _node_start(first_method)) - 1) if first_method is not None else node.end_lineno
        if first_method is None or header_end >= start:
            emit("class_header", node.name, start, max(header_end, start), [s.lineno for s in node.body if s.lineno <= header_end], doc_span(node))
        pending: list = []

        def flush_pending():
            if pending:
                emit("class_body", node.name, pending[0].lineno, pending[-1].end_lineno, [s.lineno for s in pending], None)
                pending.clear()
        for stmt in node.body:
            if isinstance(stmt, (ast.FunctionDef, ast.AsyncFunctionDef)):
                flush_pending()
                fn(stmt, f"{node.name}.{stmt.name}", "method")
            elif stmt.lineno > header_end:             # statements after the first method (the ones before it are in the header)
                pending.append(stmt)
        flush_pending()

    for node in body:
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
            flush_run()
            fn(node, node.name, "function")
        elif isinstance(node, ast.ClassDef):
            flush_run()
            class_block(node)
        elif isinstance(node, (ast.Import, ast.ImportFrom)):
            flush_run()
        elif _is_pure_literal_assign(node):
            flush_run()                            # its value is already a `python_literal` observation
        else:
            run.append(node)
    flush_run()
    return blocks, truncated[0]


MAX_GAP_RANGES = 20


def _unobserved_ranges(tree, lines: list[str], observations: list[dict]) -> list[tuple[int, int]]:
    """Non-blank lines that no observation spans and that are not import statements: the honest coverage statement of the code parser (comments after the last statement
    of a block, comment-only lines between blocks). Imports are skipped on purpose: they are dependency hints, never claimed as observed."""
    covered: set[int] = set()
    for o in observations:
        loc = o.get("locator") or {}
        if isinstance(loc.get("line_start"), int) and isinstance(loc.get("line_end"), int):
            covered.update(range(loc["line_start"], loc["line_end"] + 1))
    for node in ast.walk(tree):
        if isinstance(node, (ast.Import, ast.ImportFrom)):
            covered.update(range(node.lineno, node.end_lineno + 1))
    missing = [n for n, text in enumerate(lines, start=1) if text.strip() and n not in covered]
    out: list[tuple[int, int]] = []
    for n in missing:
        if out and out[-1][1] == n - 1:
            out[-1] = (out[-1][0], n)
        else:
            out.append((n, n))
    return out


def parse_python(logical_path: str, content: bytes) -> ParseResult:
    text, encoding, newline = _decode(content)
    name = "python.ast_literals_and_blocks.v1"
    try:
        tree = ast.parse(text, filename=logical_path)
    except SyntaxError as exc:
        return ParseResult("failed", name, encoding, newline, [],
                           [_diag("syntax_error", str(exc), line_start=exc.lineno or 1)], PYTHON_PARSER_VERSION)
    except (ValueError, RecursionError, MemoryError) as exc:     # null bytes, absurd nesting: this artifact fails, the worker does not
        return ParseResult("failed", name, encoding, newline, [],
                           [_diag("unparseable", f"{type(exc).__name__}: the source could not be parsed")], PYTHON_PARSER_VERSION)
    observations, diagnostics = [], []
    for node in tree.body:
        if not isinstance(node, ast.Assign):
            continue
        try:
            literal = ast.literal_eval(node.value)
        except (ValueError, TypeError, SyntaxError, MemoryError, RecursionError):
            continue
        end_line = getattr(node, "end_lineno", node.lineno)
        for target in node.targets:
            if isinstance(target, ast.Name):
                locator = {"type": "symbol_span", "symbol": target.id,
                          "line_start": node.lineno, "line_end": end_line}
                try:
                    observations.append(_obs("python_literal", locator, literal, version=PYTHON_PARSER_VERSION))
                except UnrepresentableValueError as exc:
                    diagnostics.append(_diag("unrepresentable_value", str(exc),
                                             symbol=target.id, line_start=node.lineno))
    if len(content) > MAX_PY_SOURCE_BYTES:
        diagnostics.append(_diag("code_blocks_skipped_source_too_large",
                                 f"source is {len(content)} bytes (limit {MAX_PY_SOURCE_BYTES}): only literal assignments are observed"))
    else:
        blocks, truncated = _code_blocks(tree, _lines(text, newline), "\r\n" if newline == "crlf" else "\n")
        for b in blocks:
            b["parser_version"] = PYTHON_PARSER_VERSION
        observations.extend(blocks)
        if truncated:
            diagnostics.append(_diag("code_blocks_truncated", f"more than {MAX_CODE_BLOCKS} code blocks: the rest are not observed"))
        gaps = _unobserved_ranges(tree, _lines(text, newline), observations)
        if gaps:
            shown = gaps[:MAX_GAP_RANGES]
            diagnostics.append(_diag("lines_not_observed",
                                     f"{sum(b - a + 1 for a, b in gaps)} non-blank line(s) in {len(gaps)} range(s) are not part of any observation (for example a comment after the last "
                                     "statement of a block): this file is not covered completely, so absence of a rule here is not evidence",
                                     ranges=[[a, b] for a, b in shown], more=len(gaps) - len(shown)))
    status = "partial" if diagnostics else "parsed"
    return ParseResult(status, name, encoding, newline, observations, diagnostics, PYTHON_PARSER_VERSION)


# ----------------------------------------------------------------- xlsx (no formula/macro execution)

def parse_xlsx(logical_path: str, content: bytes) -> ParseResult:
    import openpyxl

    try:
        wb_formula = openpyxl.load_workbook(io.BytesIO(content), data_only=False, read_only=True)
        wb_value = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # pragma: no cover - admission already validated the workbook
        return ParseResult("failed", "xlsx.cells.v1", "binary", None, [],
                           [_diag("invalid_workbook", str(exc), line_start=1)])
    observations, diagnostics = [], []
    for sheet_name in wb_formula.sheetnames:
        f_sheet, v_sheet = wb_formula[sheet_name], wb_value[sheet_name]
        for row in f_sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                locator = {"type": "cell", "sheet": sheet_name, "cell": cell.coordinate}
                try:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        cached = v_sheet[cell.coordinate].value
                        fields = {"formula_text": cell.value, "cached_value": cached,
                                  # A cached value can never be reported as CURRENT here: no
                                  # engine in this pipeline recomputes formulas (support
                                  # matrix §1, XLSX row).
                                  "recalc_status": "not_calculated" if cached is None else "unknown"}
                        observations.append(_obs("cell", locator, None, fields))
                    else:
                        observations.append(_obs("cell", locator, cell.value))
                except UnrepresentableValueError as exc:
                    diagnostics.append(_diag("unrepresentable_value", str(exc),
                                             sheet=sheet_name, cell=cell.coordinate))
    wb_formula.close()
    wb_value.close()
    status = "partial" if diagnostics else "parsed"
    return ParseResult(status, "xlsx.cells.v1", "binary", None, observations, diagnostics)


# ----------------------------------------------------------------- csv / jsonl (observations, never rules)

def _id_column(headers: list[str]) -> int:
    lowered = [h.strip().lower() for h in headers]
    if "id" in lowered:
        return lowered.index("id")
    for i, h in enumerate(lowered):
        if h.endswith("_id"):
            return i
    return 0


def parse_csv(logical_path: str, content: bytes) -> ParseResult:
    text, encoding, newline = _decode(content)
    reader = csv.reader(io.StringIO(text))
    try:
        header = next(reader)
    except StopIteration:
        return ParseResult("parsed", "csv.records.v1", encoding, newline, [], [])
    id_idx = _id_column(header)
    currency_idx = next((i for i, h in enumerate(header) if h.strip().lower() == "currency"), None)
    observations, diagnostics = [], []
    seen_ids: set[str] = set()
    currencies: set[str] = set()
    for row in reader:
        if not row:
            continue
        row = row + [""] * (len(header) - len(row))
        row_id = row[id_idx]
        fields = {header[i]: row[i] for i in range(len(header)) if i != id_idx}
        if row_id in seen_ids:
            diagnostics.append(_diag("duplicate_row_id", f"row id {row_id!r} repeats",
                                     row_id=row_id))
        seen_ids.add(row_id)
        for name, value in ([(header[id_idx], row_id)] + list(fields.items())):
            if value == "":
                diagnostics.append(_diag("missing_field", f"{name!r} is empty",
                                         row_id=row_id, field=name))
        if currency_idx is not None and row[currency_idx]:
            currencies.add(row[currency_idx])
        observations.append(_obs("record", {"type": "row", "row_id": row_id}, row_id, fields))
    if len(currencies) > 1:
        diagnostics.append(_diag("mixed_currency", f"multiple currencies present: {sorted(currencies)}"))
    return ParseResult("parsed", "csv.records.v1", encoding, newline, observations, diagnostics)


def _record_id_key(keys: list[str]) -> str:
    lowered = {k.lower(): k for k in keys}
    if "id" in lowered:
        return lowered["id"]
    for k in keys:
        if k.lower().endswith("_id"):
            return k
    return keys[0] if keys else "id"


def parse_jsonl(logical_path: str, content: bytes) -> ParseResult:
    text, encoding, newline = _decode(content)
    observations, diagnostics = [], []
    ok, bad = 0, 0
    first_keys: frozenset | None = None
    for lineno, raw_line in enumerate(_lines(text, newline), start=1):
        if not raw_line.strip():
            continue
        try:
            obj = json.loads(raw_line)
        except ValueError as exc:
            diagnostics.append(_diag("invalid_json_line", str(exc), line=lineno))
            bad += 1
            continue
        if not isinstance(obj, dict):
            diagnostics.append(_diag("invalid_json_line", "line is not a JSON object", line=lineno))
            bad += 1
            continue
        ok += 1
        keys = list(obj.keys())
        id_key = _record_id_key(keys)
        row_id = obj.get(id_key)
        fields = {k: v for k, v in obj.items() if k != id_key}
        other_keys = frozenset(fields)
        if first_keys is None:
            first_keys = other_keys
        elif other_keys != first_keys:
            diagnostics.append(_diag("schema_mismatch", "record fields differ from the first record",
                                     row_id=row_id))
        observations.append(_obs("record", {"type": "row", "row_id": row_id}, row_id, fields))
    if bad == 0:
        status = "parsed"
    elif ok == 0:
        status = "failed"
    else:
        status = "partial"
    return ParseResult(status, "jsonl.records.v1", encoding, newline, observations, diagnostics)


_PARSERS = {
    ".md": parse_markdown, ".yml": parse_yaml, ".yaml": parse_yaml,
    ".py": parse_python, ".xlsx": parse_xlsx, ".csv": parse_csv, ".jsonl": parse_jsonl,
}


def parser_for(logical_path: str):
    lower = logical_path.lower()
    for ext, fn in _PARSERS.items():
        if lower.endswith(ext):
            return fn
    return None
