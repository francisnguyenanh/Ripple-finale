"""`spreadsheet_rules` doc adapter — where rules actually live.

Excel is where business rules really are kept, and where drift happens most
often: it is the one artifact a business owner can change without asking
anyone. The approval matrix in this seed was lowered in June 2022 and no
other representation moved.

Reads from bytes so a snapshot -- baseline or overlay -- can be read without
anything touching disk. Emits `spreadsheet_claim`, deliberately NOT `rule`:
a spreadsheet says what it says, and what is TRUE is settled by comparing
representations, never by trusting one (I9).

Imports nothing from `core.kernel`.
"""

from __future__ import annotations

import io
from typing import Iterator

from openpyxl import load_workbook

from core.adapters.base import IRRecord, Provenance

#: Column headers that name a rule's threshold, in the spellings a finance
#: team actually uses.
_VALUE_HEADERS = {"threshold", "limit", "amount", "value", "ceiling"}
_ID_HEADERS = {"rule id", "rule_id", "id"}
_NAME_HEADERS = {"rule name", "rule_name", "name", "description"}
_UNIT_HEADERS = {"unit", "currency"}
_REVIEWED_HEADERS = {"last reviewed", "last_reviewed", "reviewed", "last assessed"}


def _index(headers: list[str], wanted: set[str]) -> int | None:
    for i, header in enumerate(headers):
        if (header or "").strip().lower() in wanted:
            return i
    return None


def extract_bytes(raw: bytes, path: str, snapshot_id: str) -> Iterator[IRRecord]:
    workbook = load_workbook(io.BytesIO(raw), data_only=True)

    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue

        headers = [str(c) if c is not None else "" for c in rows[0]]
        value_at = _index(headers, _VALUE_HEADERS)
        if value_at is None:
            continue

        id_at = _index(headers, _ID_HEADERS)
        name_at = _index(headers, _NAME_HEADERS)
        unit_at = _index(headers, _UNIT_HEADERS)
        reviewed_at = _index(headers, _REVIEWED_HEADERS)

        for offset, row in enumerate(rows[1:], start=2):
            value = row[value_at] if value_at < len(row) else None
            if not isinstance(value, (int, float)) or isinstance(value, bool):
                continue

            rule_id = str(row[id_at]) if id_at is not None and id_at < len(row) and row[id_at] else ""
            rule_name = str(row[name_at]) if name_at is not None and name_at < len(row) and row[name_at] else ""
            unit = str(row[unit_at]) if unit_at is not None and unit_at < len(row) and row[unit_at] else ""
            reviewed = str(row[reviewed_at]) if reviewed_at is not None and reviewed_at < len(row) and row[reviewed_at] else ""

            cell = f"{sheet.cell(row=offset, column=value_at + 1).column_letter}{offset}"
            yield IRRecord(
                kind="entity",
                id=f"xlsx:{path}:{sheet_name}:{cell}",
                payload={
                    "type": "spreadsheet_claim",
                    "rule_id": rule_id,
                    "rule_name": rule_name,
                    "value": int(value),
                    "unit": unit,
                    # A threshold table states a lower bound even when no
                    # operator column exists; saying so explicitly is what
                    # lets tier 2 compare it with the workflow's `>=`.
                    "field": "amount",
                    "op": ">=",
                    "sheet": sheet_name,
                    "cell": cell,
                    "last_reviewed": reviewed,
                    "semantic_hint": "",
                },
                provenance=(Provenance(doc=path, sheet=sheet_name, row=offset,
                                       method="parser",
                                       # The row as it really reads, not a
                                       # sentence assembled from it. The
                                       # assembled form ("BR-001
                                       # high_value_requires_manager =
                                       # 500000 JPY") appears nowhere in the
                                       # workbook, so it cannot be checked
                                       # against the artifact -- and by I4
                                       # an unverifiable quote is not
                                       # evidence, however true it is.
                                       #
                                       # This is the 500,000 side of the
                                       # hero conflict. The one claim the
                                       # product most needs to stand behind
                                       # was the one it could not.
                                       quote=" | ".join(
                                           str(c) for c in row if c is not None)),))
